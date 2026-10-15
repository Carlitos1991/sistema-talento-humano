import calendar
import datetime as dt
from datetime import date as dt_date
from io import BytesIO

import openpyxl
from dateutil.relativedelta import relativedelta
from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin, PermissionRequiredMixin
from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger
from django.db import IntegrityError, transaction
from django.db.models import Prefetch, Count
from django.db.models import Q, F
from django.http import HttpResponse
from django.http import JsonResponse
from django.shortcuts import get_object_or_404
from django.shortcuts import redirect, render
from django.template.loader import get_template
from django.template.loader import render_to_string
from django.urls import reverse_lazy, reverse
from django.views import View
from django.views.generic import ListView, TemplateView, CreateView, UpdateView, FormView
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from xhtml2pdf import pisa

from budget.models import BudgetLine
from contract.models import ManagementPeriod
from core.models import SystemConfiguration
from permitrequest.models import PermitType
from person.models import Person
from personnel_actions.models import PersonnelAction, ActionType
from .forms import *
from .models import *
from .services import (
    calculate_available_days_for_month, get_vacation_regime_info
)
from .services import resolve_scheduling_hierarchy

MONTH_NAMES = {
    1: 'Enero', 2: 'Febrero', 3: 'Marzo', 4: 'Abril',
    5: 'Mayo', 6: 'Junio', 7: 'Julio', 8: 'Agosto',
    9: 'Septiembre', 10: 'Octubre', 11: 'Noviembre', 12: 'Diciembre'
}


# ==============================================================================
# CALCULO POR REGIMEN
# ==============================================================================

def calculate_earned_days(employee, calculation_date=None):
    """
    Calcula los días ganados de vacaciones:
    - LOSEP / Empleado: 30 días fijos.
    - TRABAJADOR (Código del Trabajo): 15 días base.
      A partir del 5to año cumplido (año 6 en adelante), 1 día adicional por año (tope 15 extra = 30 máx).
    """
    if not employee or not employee.employment_status:
        return Decimal('30.0')

    regime = (employee.employment_status.code or '').upper()

    # Régimen Código del Trabajo / Trabajador
    if 'TRABAJADOR' in regime or 'CÓDIGO' in regime or 'CODIGO' in regime:
        base_days = Decimal('15.0')
        extra_days = Decimal('0.0')

        # Contar periodos previos asignados
        previous_periods_count = EmployeeVacationBalance.objects.filter(employee=employee).count()

        # Calcular años de servicio si existe fecha de ingreso formal
        entry_date = getattr(employee, 'entry_date', None) or getattr(employee.person, 'entry_date', None)
        if entry_date:
            anniversary_date = calculation_date or timezone.now().date()
            years_of_service = relativedelta(anniversary_date, entry_date).years
        else:
            years_of_service = previous_periods_count

        # A partir del quinto año cumplido (desde el año 6)
        if years_of_service >= 5:
            bonus = min(years_of_service - 4, 15)
            extra_days = Decimal(str(bonus))

        return base_days + extra_days

    # Por defecto todo empleado / servidor público recibe 30 días
    return Decimal('30.0')


# ==============================================================================
# LISTA DE VACACIONES
# ==============================================================================

class VacationRequestListView(LoginRequiredMixin, ListView):
    """
    Listado principal de empleados para administrar vacaciones.
    """
    model = Employee
    template_name = 'vacation/vacation_request_list.html'
    context_object_name = 'employees_data'
    paginate_by = 10

    def get_queryset(self):
        qs = Employee.objects.filter(
            is_active=True
        ).select_related(
            'person',
            'area'
        ).prefetch_related(
            Prefetch(
                'employeevacationbalance_set',
                queryset=EmployeeVacationBalance.objects.select_related('period').filter(is_active=True)
            )
        ).order_by('person__last_name', 'person__first_name')

        q = self.request.GET.get('q', '').strip()
        if q:
            terms = q.split()
            for term in terms:
                qs = qs.filter(
                    Q(person__first_name__icontains=term) |
                    Q(person__last_name__icontains=term) |
                    Q(person__document_number__icontains=term)
                )

        return qs

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        employees_data = []

        for employee in context['object_list']:
            budget = BudgetLine.objects.filter(
                current_employee=employee,
                is_active=True
            ).select_related('position_item').first()

            vacation_balance = employee.employeevacationbalance_set.filter(is_active=True).first()

            # Verificación en disco de la foto física
            has_valid_photo = False
            person = employee.person
            if person and person.photo:
                try:
                    has_valid_photo = person.photo.storage.exists(person.photo.name)
                except Exception:
                    has_valid_photo = False

            employees_data.append({
                'employee': employee,
                'budget': budget,
                'vacation_balance': vacation_balance,
                'has_vacation': vacation_balance is not None,
                'has_valid_photo': has_valid_photo
            })

        context['employees_data'] = employees_data
        context['search_query'] = self.request.GET.get('q', '')
        return context

    def render_to_response(self, context, **response_kwargs):
        if self.request.headers.get('x-requested-with') == 'XMLHttpRequest':
            html = render_to_string(
                'vacation/partials/partial_vacation_employee_list.html',
                context,
                request=self.request
            )
            page_obj = context.get('page_obj')
            pagination_data = {
                'start_index': page_obj.start_index() if page_obj else 0,
                'end_index': page_obj.end_index() if page_obj else 0,
                'total_count': page_obj.paginator.count if page_obj else 0,
                'current_page': page_obj.number if page_obj else 1,
                'total_pages': page_obj.paginator.num_pages if page_obj else 1,
                'has_previous': page_obj.has_previous() if page_obj else False,
                'has_next': page_obj.has_next() if page_obj else False,
            }
            return JsonResponse({'html': html, 'pagination': pagination_data})

        return super().render_to_response(context, **response_kwargs)


class VacationCreateView(LoginRequiredMixin, TemplateView):
    template_name = 'vacation/modals/modal_vacation_form.html'


# ==============================================================================
# ADMINISTRACION DE VACACIONES
# ==============================================================================

class PeriodListView(LoginRequiredMixin, ListView):
    model = VacationPeriod
    template_name = 'vacation/period_list.html'
    context_object_name = 'periods'
    paginate_by = 10

    def get_queryset(self):
        show_inactive = self.request.GET.get('show_inactive') == 'true'

        # Por defecto solo activos
        if show_inactive:
            qs = VacationPeriod.objects.all()
        else:
            qs = VacationPeriod.objects.filter(is_active=True)

        q = self.request.GET.get('q', '').strip()
        if q:
            qs = qs.filter(name__icontains=q)

        return qs.order_by('-name')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['search_query'] = self.request.GET.get('q', '')
        context['show_inactive'] = self.request.GET.get('show_inactive') == 'true'
        return context

    def render_to_response(self, context, **response_kwargs):
        if self.request.headers.get('x-requested-with') == 'XMLHttpRequest':
            html = render_to_string(
                'vacation/partials/partial_period_list.html',
                context,
                request=self.request
            )
            return HttpResponse(html)
        return super().render_to_response(context, **response_kwargs)


class PeriodCreateView(LoginRequiredMixin, CreateView):
    model = VacationPeriod
    form_class = PeriodForm
    template_name = 'vacation/modals/modal_period_vacation_form.html'
    success_url = reverse_lazy('vacation:period_list')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['titulo'] = 'Crear Nuevo Periodo'
        context['action_url'] = reverse_lazy('vacation:period_create')
        return context

    def form_valid(self, form):
        try:
            self.object = form.save()
            if self.request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({
                    'status': 'success',
                    'success': True,
                    'message': f'Periodo "{self.object.name}" creado exitosamente',
                    'redirect_url': str(self.success_url)
                })
            messages.success(self.request, f'Periodo "{self.object.name}" creado exitosamente')
            return redirect(self.success_url)
        except IntegrityError:
            if self.request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({
                    'status': 'error',
                    'success': False,
                    'message': 'Ya existe un periodo con este nombre',
                    'errors': {'name': 'Ya existe un periodo con este nombre'}
                }, status=400)
            form.add_error('name', 'Ya existe un periodo con este nombre')
            return self.form_invalid(form)

    def form_invalid(self, form):
        if self.request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            errors = {k: v[0] for k, v in form.errors.items()}
            return JsonResponse({
                'status': 'error',
                'success': False,
                'message': 'Por favor corrija los errores en el formulario',
                'errors': errors
            }, status=400)
        return super().form_invalid(form)


class PeriodUpdateView(LoginRequiredMixin, UpdateView):
    model = VacationPeriod
    form_class = PeriodForm
    template_name = 'vacation/modals/modal_period_vacation_form.html'
    success_url = reverse_lazy('vacation:period_list')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['titulo'] = 'Editar Periodo'
        context['action_url'] = reverse_lazy('vacation:period_edit', kwargs={'pk': self.object.pk})
        context['is_edit'] = True
        return context

    def form_valid(self, form):
        try:
            self.object = form.save()
            if self.request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({
                    'status': 'success',
                    'success': True,
                    'message': f'Periodo "{self.object.name}" actualizado exitosamente',
                    'redirect_url': str(self.success_url)
                })
            messages.success(self.request, f'Periodo "{self.object.name}" actualizado exitosamente')
            return redirect(self.success_url)
        except IntegrityError:
            if self.request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({
                    'status': 'error',
                    'success': False,
                    'message': 'Ya existe un periodo con este nombre',
                    'errors': {'name': 'Ya existe un periodo con este nombre'}
                }, status=400)
            form.add_error('name', 'Ya existe un periodo con este nombre')
            return self.form_invalid(form)

    def form_invalid(self, form):
        if self.request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            errors = {k: v[0] for k, v in form.errors.items()}
            return JsonResponse({
                'status': 'error',
                'success': False,
                'message': 'Por favor corrija los errores en el formulario',
                'errors': errors
            }, status=400)
        return super().form_invalid(form)


# ==============================================================================
# PRIMER PERIODO DE VACACIONES
# ==============================================================================

class CreateFirstVacationView(LoginRequiredMixin, CreateView):
    model = EmployeeVacationBalance
    form_class = FirstVacationForm
    template_name = 'vacation/modals/modal_first_vacation_form.html'

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        employee_id = self.kwargs.get('employee_id')
        kwargs['employee_id'] = employee_id

        employee = get_object_or_404(
            Employee.objects.select_related('employment_status', 'person'),
            pk=employee_id
        )
        kwargs['initial_days'] = calculate_earned_days(employee)
        return kwargs

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        employee = get_object_or_404(
            Employee.objects.select_related('person', 'employment_status').prefetch_related(
                'current_budget_line__position_item'),
            pk=self.kwargs['employee_id']
        )

        has_valid_photo = False
        if employee.person and employee.person.photo:
            try:
                has_valid_photo = employee.person.photo.storage.exists(employee.person.photo.name)
            except Exception:
                has_valid_photo = False

        context['employee'] = employee
        context['has_valid_photo'] = has_valid_photo
        context['titulo'] = f'Crear Primera Vacación para {employee.person.full_name}'
        context['action_url'] = reverse('vacation:create_first_vacation', kwargs={'employee_id': employee.id})
        return context

    def form_valid(self, form):
        employee = get_object_or_404(Employee, pk=self.kwargs['employee_id'])

        days = form.cleaned_data.get('total_days', Decimal('0'))
        hours = form.cleaned_data.get('hours', 0)
        minutes = form.cleaned_data.get('minutes', 0)
        user_detail = form.cleaned_data.get('observation_detail', '')

        extra_from_hours = Decimal(str(hours)) * FACTOR_HOUR
        extra_from_minutes = Decimal(str(minutes)) * FACTOR_MINUTE
        final_calculated_days = days + extra_from_hours + extra_from_minutes

        balance = form.save(commit=False)
        balance.employee = employee
        balance.total_days = final_calculated_days
        balance.balance_days = final_calculated_days
        balance.additional_days = Decimal('0.0')

        period_name = form.cleaned_data['period'].name
        balance.observation = f"CARGA INICIAL PERIODO {period_name}: {user_detail}"
        balance.is_active = True
        balance.created_by = self.request.user
        balance.save()

        redirect_url = reverse('vacation:employee_vacation_detail', kwargs={'employee_id': employee.id})
        if self.request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return JsonResponse({
                'status': 'success',
                'success': True,
                'message': 'Periodo inicial creado exitosamente.',
                'redirect_url': redirect_url
            })
        return redirect(redirect_url)

    def form_invalid(self, form):
        if self.request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            errors = {k: v[0] for k, v in form.errors.items()}
            return JsonResponse({
                'status': 'error',
                'success': False,
                'message': 'Por favor corrija los errores en el formulario',
                'errors': errors
            }, status=400)
        return super().form_invalid(form)


class CreateNewVacationPeriodView(LoginRequiredMixin, CreateView):
    model = EmployeeVacationBalance
    form_class = FirstVacationForm
    template_name = 'vacation/modals/modal_new_vacation_form.html'

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        employee_id = self.kwargs.get('employee_id')
        kwargs['employee_id'] = employee_id

        employee = get_object_or_404(Employee.objects.select_related('employment_status', 'person'), pk=employee_id)
        kwargs['initial_days'] = calculate_earned_days(employee)
        return kwargs

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        employee = get_object_or_404(
            Employee.objects.select_related('person', 'employment_status').prefetch_related(
                'current_budget_line__position_item'),
            pk=self.kwargs['employee_id']
        )

        has_valid_photo = False
        if employee.person and employee.person.photo:
            try:
                has_valid_photo = employee.person.photo.storage.exists(employee.person.photo.name)
            except Exception:
                has_valid_photo = False

        context['employee'] = employee
        context['has_valid_photo'] = has_valid_photo
        context['titulo'] = f'Nueva Vacación para {employee.person.full_name}'
        context['action_url'] = reverse('vacation:create_new_vacation', kwargs={'employee_id': employee.id})
        return context

    def form_valid(self, form):
        employee = get_object_or_404(Employee.objects.select_related('employment_status'),
                                     pk=self.kwargs['employee_id'])
        previous_balances = EmployeeVacationBalance.objects.filter(employee=employee).order_by('created_at')

        is_trabajador = employee.employment_status and 'TRABAJADOR' in (employee.employment_status.code or '').upper()
        new_period_days = calculate_earned_days(employee)

        last_balance = previous_balances.last()
        previous_balance = last_balance.balance_days if last_balance else Decimal('0.0')
        calculated_balance = previous_balance + new_period_days
        max_limit = Decimal('45.0') if is_trabajador else Decimal('60.0')

        lost_days = Decimal('0.0')
        if calculated_balance > max_limit:
            lost_days = calculated_balance - max_limit
            final_balance = max_limit
        else:
            final_balance = calculated_balance

        period_name = form.cleaned_data['period'].name
        if lost_days > 0:
            observation = (
                f"Se creó el período {period_name} con un balance de {final_balance} días. "
                f"IMPORTANTE: El empleado perdió {lost_days} días del período {period_name} "
                f"por exceder el límite máximo legal "
                f"(balance anterior: {previous_balance} días + {new_period_days} días nuevos = {calculated_balance} días)."
            )
        else:
            observation = (
                f"Se creó el período {period_name} con un balance de {final_balance} días "
                f"(balance anterior: {previous_balance} días + {new_period_days} días nuevos)."
            )

        balance = form.save(commit=False)
        balance.employee = employee
        balance.total_days = new_period_days
        balance.balance_days = final_balance
        balance.observation = observation
        balance.additional_days = last_balance.balance_days if last_balance else Decimal('0.0')
        balance.is_active = True
        balance.created_by = self.request.user
        balance.save()

        redirect_url = reverse('vacation:employee_vacation_detail', kwargs={'employee_id': employee.id})
        if self.request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return JsonResponse({
                'status': 'success',
                'success': True,
                'message': 'Nuevo periodo de vacaciones creado exitosamente.',
                'redirect_url': redirect_url
            })
        messages.success(self.request, 'Nuevo periodo de vacaciones creado exitosamente')
        return redirect(redirect_url)

    def form_invalid(self, form):
        if self.request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            errors = {k: v[0] for k, v in form.errors.items()}
            return JsonResponse({
                'status': 'error',
                'success': False,
                'message': 'Por favor corrija los errores en el formulario',
                'errors': errors
            }, status=400)
        return super().form_invalid(form)


# ==============================================================================
# DETALLE DE VACACIONES
# ==============================================================================

class EmployeeVacationDetailView(LoginRequiredMixin, TemplateView):
    template_name = 'vacation/employee_vacation_detail.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        employee = get_object_or_404(
            Employee.objects.select_related('person', 'employment_status', 'area'),
            pk=self.kwargs['employee_id']
        )

        search_query = self.request.GET.get('search', '').strip()
        page_number = self.request.GET.get('page', 1)
        per_page = int(self.request.GET.get('per_page', 10))

        active_balance = EmployeeVacationBalance.objects.filter(
            employee=employee,
            is_active=True,
            period__is_active=True
        ).select_related('period').first()

        vacation_balances = EmployeeVacationBalance.objects.filter(
            employee=employee,
            is_active=True
        ).select_related('period')

        if search_query:
            vacation_balances = vacation_balances.filter(period__name__icontains=search_query)

        vacation_balances = vacation_balances.order_by('-created_at')
        paginator = Paginator(vacation_balances, per_page)

        try:
            vacation_balances_page = paginator.page(page_number)
        except (PageNotAnInteger, EmptyPage):
            vacation_balances_page = paginator.page(1)

        context['employee'] = employee
        context['active_balance'] = active_balance
        context['vacation_balances'] = vacation_balances_page
        context['search_query'] = search_query
        context['per_page'] = per_page
        return context


# ==============================================================================
# VACATION CHARGED PERMITS
# ==============================================================================

class CreateHourPermitVacationView(LoginRequiredMixin, FormView):
    template_name = 'vacation/modals/modal_hour_permit_form.html'
    form_class = HourPermitVacationForm

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        employee = get_object_or_404(Employee, pk=self.kwargs['employee_id'])
        active_balance = EmployeeVacationBalance.objects.filter(
            employee=employee,
            is_active=True,
            period__is_active=True
        ).select_related('period').first()

        context['employee'] = employee
        context['active_balance'] = active_balance
        context['titulo'] = f'Permiso por Horas - {employee.person.full_name}'
        context['action_url'] = reverse('vacation:create_hour_permit', kwargs={'employee_id': employee.id})
        return context

    def form_valid(self, form):
        try:
            employee = get_object_or_404(Employee, pk=self.kwargs['employee_id'])
            active_balance = EmployeeVacationBalance.objects.filter(
                employee=employee,
                is_active=True,
                period__is_active=True
            ).select_related('period').first()

            if not active_balance:
                return JsonResponse({
                    'status': 'error',
                    'success': False,
                    'message': 'El empleado no tiene un periodo de vacaciones activo'
                }, status=400)

            start_date = form.cleaned_data['start_date']
            start_time = form.cleaned_data['start_time']
            hours = form.cleaned_data.get('hours', 0) or 0
            minutes = form.cleaned_data.get('minutes', 0) or 0

            from datetime import datetime, timedelta
            total_minutes_permit = (hours * 60) + minutes
            end_time_calculated = (
                    datetime.combine(start_date, start_time) + timedelta(minutes=total_minutes_permit)
            ).time()

            existing_permit = PermitRequest.objects.filter(
                employee=employee,
                start_date=start_date,
                status__in=['REQUESTED', 'APPROVED']
            ).exclude(start_time__gte=end_time_calculated).exclude(end_time__lte=start_time).first()

            if existing_permit:
                return JsonResponse({
                    'status': 'error',
                    'success': False,
                    'message': 'Ya existe un permiso registrado en ese horario.'
                }, status=400)

            factor_hour_base = Decimal('0.125')
            factor_minute_base = Decimal('0.00208333')
            proportional_hour = Decimal('0.05')
            proportional_minute = Decimal('0.00083')

            value_discount = (Decimal(str(hours)) * factor_hour_base) + (Decimal(str(minutes)) * factor_minute_base)
            proportional_discount = (Decimal(str(hours)) * proportional_hour) + (
                    Decimal(str(minutes)) * proportional_minute)
            total_discount = value_discount + proportional_discount

            if active_balance.balance_days < total_discount:
                return JsonResponse({
                    'status': 'error',
                    'success': False,
                    'message': f'Saldo insuficiente. Disponible: {active_balance.balance_days} días, Requerido: {total_discount:.2f} días'
                }, status=400)

            permit_type = PermitType.objects.filter(name='Personales', is_active=True).first()
            if not permit_type:
                return JsonResponse({
                    'status': 'error',
                    'success': False,
                    'message': 'No se encontró el tipo de permiso "Personales" activo'
                }, status=400)

            end_time = (datetime.combine(start_date, start_time) + timedelta(minutes=total_minutes_permit)).time()

            PermitRequest.objects.create(
                employee=employee,
                permit_type=permit_type,
                start_date=start_date,
                start_time=start_time,
                end_date=start_date,
                end_time=end_time,
                days=0,
                hours=hours,
                minutes=minutes,
                status='REQUESTED',
                created_by=self.request.user
            )

            return JsonResponse({
                'status': 'success',
                'success': True,
                'message': 'Solicitud de permiso creada exitosamente. Pendiente de aprobación.',
                'redirect_url': reverse('vacation:employee_vacation_detail', kwargs={'employee_id': employee.id})
            })
        except Exception as e:
            return JsonResponse({'status': 'error', 'success': False, 'message': str(e)}, status=500)

    def form_invalid(self, form):
        errors = {k: v[0] for k, v in form.errors.items()}
        return JsonResponse({
            'status': 'error',
            'success': False,
            'message': 'Por favor corrija los errores en el formulario',
            'errors': errors
        }, status=400)


class CreateDayPermitVacationView(LoginRequiredMixin, FormView):
    template_name = 'vacation/modals/modal_day_permit_form.html'
    form_class = DayPermitVacationForm

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        employee = get_object_or_404(Employee, pk=self.kwargs['employee_id'])
        active_balance = EmployeeVacationBalance.objects.filter(
            employee=employee,
            is_active=True,
            period__is_active=True
        ).select_related('period').first()

        context['employee'] = employee
        context['active_balance'] = active_balance
        context['titulo'] = f'Permiso por Días - {employee.person.full_name}'
        context['action_url'] = reverse('vacation:create_day_permit', kwargs={'employee_id': employee.id})
        return context

    def form_valid(self, form):
        try:
            employee = get_object_or_404(Employee, pk=self.kwargs['employee_id'])
            active_balance = EmployeeVacationBalance.objects.filter(
                employee=employee,
                is_active=True,
                period__is_active=True
            ).select_related('period').first()

            if not active_balance:
                return JsonResponse({
                    'status': 'error',
                    'success': False,
                    'message': 'El empleado no tiene un periodo de vacaciones activo'
                }, status=400)

            start_date = form.cleaned_data['start_date']
            days = form.cleaned_data['days']

            from datetime import timedelta, time
            end_date_calculated = start_date + timedelta(days=days - 1)

            existing_permit = PermitRequest.objects.filter(
                employee=employee,
                status__in=['REQUESTED', 'APPROVED'],
                start_date__lte=end_date_calculated,
                end_date__gte=start_date
            ).first()

            if existing_permit:
                return JsonResponse({
                    'status': 'error',
                    'success': False,
                    'message': 'Ya existe un permiso registrado en estas fechas.'
                }, status=400)

            factor_day = Decimal('1.0')
            proportional_day = Decimal('0.4')
            total_discount = (Decimal(str(days)) * factor_day) + (Decimal(str(days)) * proportional_day)

            if active_balance.balance_days < total_discount:
                return JsonResponse({
                    'status': 'error',
                    'success': False,
                    'message': f'Saldo insuficiente. Disponible: {active_balance.balance_days} días, Requerido: {total_discount:.2f} días'
                }, status=400)

            permit_type = PermitType.objects.filter(name='Personales', is_active=True).first()
            if not permit_type:
                return JsonResponse({
                    'status': 'error',
                    'success': False,
                    'message': 'No se encontró el tipo de permiso "Personales"'
                }, status=400)

            PermitRequest.objects.create(
                employee=employee,
                permit_type=permit_type,
                start_date=start_date,
                start_time=time(8, 0),
                end_date=end_date_calculated,
                end_time=None,
                days=days,
                hours=0,
                minutes=0,
                status='REQUESTED',
                created_by=self.request.user
            )

            return JsonResponse({
                'status': 'success',
                'success': True,
                'message': 'Solicitud de permiso creada exitosamente. Pendiente de aprobación.',
                'redirect_url': reverse('vacation:employee_vacation_detail', kwargs={'employee_id': employee.id})
            })
        except Exception as e:
            return JsonResponse({'status': 'error', 'success': False, 'message': str(e)}, status=500)

    def form_invalid(self, form):
        errors = {k: v[0] for k, v in form.errors.items()}
        return JsonResponse({
            'status': 'error',
            'success': False,
            'message': 'Por favor corrija los errores en el formulario',
            'errors': errors
        }, status=400)


class EmployeePermitListView(LoginRequiredMixin, TemplateView):
    template_name = 'vacation/modals/modal_permit_list.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        employee = get_object_or_404(Employee.objects.select_related('person'), pk=self.kwargs['employee_id'])

        try:
            personal_permit_type = PermitType.objects.get(name='Personales', is_active=True)
            permits_queryset = PermitRequest.objects.filter(
                employee=employee,
                permit_type=personal_permit_type
            ).select_related('permit_type', 'employee__person')

            search_query = self.request.GET.get('search', '').strip()
            if search_query:
                from datetime import datetime
                try:
                    search_date = datetime.strptime(search_query, '%d/%m/%Y').date()
                    permits_queryset = permits_queryset.filter(Q(start_date=search_date) | Q(end_date=search_date))
                except ValueError:
                    pass
                context['search_query'] = search_query

            permits_queryset = permits_queryset.order_by('-start_date', '-id')
            paginator = Paginator(permits_queryset, 10)
            permits = paginator.page(self.request.GET.get('page', 1))
        except Exception:
            permits = []

        context['employee'] = employee
        context['permits'] = permits
        context['titulo'] = f'Listado de Permisos - {employee.person.full_name}'
        return context


class ApprovePermitView(LoginRequiredMixin, View):
    def post(self, request, permit_id):
        try:
            permit = get_object_or_404(PermitRequest, pk=permit_id)
            if permit.status != 'REQUESTED':
                return JsonResponse({
                    'status': 'error',
                    'success': False,
                    'message': f'Estado actual: {permit.get_status_display()}'
                }, status=400)

            active_balance = EmployeeVacationBalance.objects.filter(
                employee=permit.employee,
                is_active=True,
                period__is_active=True
            ).select_related('period').first()

            if not active_balance:
                return JsonResponse({'status': 'error', 'success': False, 'message': 'No tiene período activo'},
                                    status=400)

            if permit.days > 0:
                value_discount = Decimal(str(permit.days)) * Decimal('1.0')
                proportional_discount = Decimal(str(permit.days)) * Decimal('0.4')
            else:
                hours = permit.hours or 0
                minutes = permit.minutes or 0
                value_discount = (Decimal(str(hours)) * Decimal('0.125')) + (
                        Decimal(str(minutes)) * Decimal('0.00208333'))
                proportional_discount = (Decimal(str(hours)) * Decimal('0.05')) + (
                        Decimal(str(minutes)) * Decimal('0.00083333'))

            total_discount = value_discount + proportional_discount
            if active_balance.balance_days < total_discount:
                return JsonResponse({'status': 'error', 'success': False, 'message': 'Saldo insuficiente.'}, status=400)

            obs = f'Permiso {permit.days}d' if permit.days > 0 else f'Permiso {permit.hours}h {permit.minutes}m'
            VacationHistory.objects.create(
                vacation_balance=active_balance,
                permit_request=permit,
                value_discount=float(value_discount),
                proportional_discount=float(proportional_discount),
                days_discount=float(permit.days) if permit.days > 0 else None,
                hours_discount=float(permit.hours) if permit.hours > 0 else None,
                minutes_discount=float(permit.minutes) if permit.minutes > 0 else None,
                observation=f"{obs} APROBADO período {active_balance.period.name}",
                created_by=request.user
            )

            EmployeeVacationBalance.objects.filter(id=active_balance.id).update(
                balance_days=F('balance_days') - total_discount,
                permit_days=F('permit_days') + total_discount
            )

            permit.status = 'APPROVED'
            permit.response_by = request.user
            permit.response_date = timezone.now()
            permit.save()

            return JsonResponse({
                'status': 'success',
                'success': True,
                'message': f'Permiso aprobado exitosamente. Descuento: {total_discount:.2f} días'
            })
        except Exception as e:
            return JsonResponse({'status': 'error', 'success': False, 'message': str(e)}, status=500)


class RejectPermitView(LoginRequiredMixin, View):
    def post(self, request, permit_id):
        permit = get_object_or_404(PermitRequest, pk=permit_id)
        if permit.status != 'REQUESTED':
            return JsonResponse({'status': 'error', 'success': False, 'message': 'Permiso ya procesado'}, status=400)

        permit.status = 'REJECTED'
        permit.response_by = request.user
        permit.response_date = timezone.now()
        permit.save()
        return JsonResponse({'status': 'success', 'success': True, 'message': 'Permiso rechazado exitosamente'})


class CancelPermitView(LoginRequiredMixin, View):
    def post(self, request, permit_id):
        permit = get_object_or_404(PermitRequest, pk=permit_id)
        if permit.status not in ['REQUESTED', 'APPROVED']:
            return JsonResponse({'status': 'error', 'success': False, 'message': 'No se puede anular este permiso'},
                                status=400)

        if permit.status == 'APPROVED':
            try:
                hist = VacationHistory.objects.get(permit_request=permit)
                total_discount = Decimal(str(hist.value_discount)) + Decimal(str(hist.proportional_discount))
                EmployeeVacationBalance.objects.filter(id=hist.vacation_balance.id).update(
                    balance_days=F('balance_days') + total_discount,
                    permit_days=F('permit_days') - total_discount
                )
                hist.delete()
            except VacationHistory.DoesNotExist:
                pass

        permit.delete()
        return JsonResponse({'status': 'success', 'success': True, 'message': 'Permiso anulado exitosamente'})


# ==============================================================================
# LIQUIDACION DE VACACIONES
# ==============================================================================

class CreateVacationLiquidationView(LoginRequiredMixin, FormView):
    template_name = 'vacation/modals/modal_liquidation_form.html'
    form_class = VacationLiquidationForm

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        employee = get_object_or_404(Employee, pk=self.kwargs['employee_id'])
        balance = EmployeeVacationBalance.objects.filter(employee=employee, is_active=True).order_by(
            '-created_at').first()
        context['employee'] = employee
        context['balance'] = balance
        context['available_days'] = balance.balance_days if balance else 0
        context['action_url'] = reverse('vacation:create_liquidation', kwargs={'employee_id': employee.id})
        return context

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        employee = get_object_or_404(Employee, pk=self.kwargs['employee_id'])
        balance = EmployeeVacationBalance.objects.filter(employee=employee, is_active=True).order_by(
            '-created_at').first()
        kwargs['available_days'] = balance.balance_days if balance else 0
        return kwargs

    def form_valid(self, form):
        try:
            employee = get_object_or_404(Employee, pk=self.kwargs['employee_id'])
            balance = EmployeeVacationBalance.objects.filter(employee=employee, is_active=True).order_by(
                '-created_at').first()

            if not balance:
                return JsonResponse({'status': 'error', 'success': False, 'message': 'Sin saldo de vacaciones activo.'},
                                    status=400)

            start_date = form.cleaned_data['start_date']
            end_date = form.cleaned_data['end_date']
            days_requested = form.cleaned_data['days_requested']

            overlapping = VacationRequest.objects.filter(
                employee=employee,
                balance_used=balance,
                status__in=['PENDING', 'APPROVED'],
                start_date__lte=end_date,
                end_date__gte=start_date
            ).exists()

            if overlapping:
                return JsonResponse(
                    {'status': 'error', 'success': False, 'message': 'Ya existe una solicitud en ese rango de fechas.'},
                    status=400)

            with transaction.atomic():
                action_type = ActionType.objects.filter(name__iexact='VACACIONES').first()
                if not action_type:
                    return JsonResponse(
                        {'status': 'error', 'success': False, 'message': 'Tipo de acción VACACIONES no encontrado.'},
                        status=400)

                current_year = dt.date.today().year
                last_action = PersonnelAction.objects.filter(number__endswith=f'-{current_year}').order_by(
                    '-number').first()
                new_number = 1
                if last_action:
                    try:
                        new_number = int(last_action.number.split('-')[0]) + 1
                    except (ValueError, IndexError):
                        new_number = 1
                action_number = f"{new_number:04d}-{current_year}"

                reintegro_date = end_date + dt.timedelta(days=1)
                explanation = (
                    f'LIQUIDACIÓN DE {days_requested} DÍAS DE VACACIONES AL SERVIDOR DESDE EL "{start_date}" '
                    f'AL "{end_date}" CORRESPONDIENTE AL PERIODO "{balance.period}". '
                    f'REINTEGRO: {reintegro_date}'
                )

                personnel_action = PersonnelAction.objects.create(
                    employee=employee,
                    action_type=action_type,
                    number=action_number,
                    explanation=explanation,
                    motivation='SOLICITUD DE VACACIONES',
                    date_issue=dt.date.today(),
                    date_effective=start_date,
                    is_registered=False,
                    authority_1=form.cleaned_data['nominating_authority'],
                    authority_2=form.cleaned_data['human_resources_responsible'],
                    register=form.cleaned_data['registration_responsible'],
                    reviewer=form.cleaned_data['review_responsible'],
                    elaboration=form.cleaned_data['elaborated_by'],
                    created_by=self.request.user
                )

                VacationRequest.objects.create(
                    employee=employee,
                    balance_used=balance,
                    start_date=start_date,
                    end_date=end_date,
                    days_quantity=Decimal(str(days_requested)),
                    status='PENDING',
                    personnel_action=personnel_action,
                    created_by=self.request.user,
                    date_issued=dt.date.today()
                )

            return JsonResponse({
                'status': 'success',
                'success': True,
                'message': f'Solicitud de liquidación creada exitosamente. Acción No. {action_number}'
            })
        except Exception as e:
            return JsonResponse({'status': 'error', 'success': False, 'message': str(e)}, status=500)

    def form_invalid(self, form):
        errors = {k: v[0] for k, v in form.errors.items()}
        return JsonResponse({
            'status': 'error',
            'success': False,
            'message': 'Por favor corrija los errores en el formulario',
            'errors': errors
        }, status=400)


class EmployeeLiquidationListView(LoginRequiredMixin, TemplateView):
    template_name = 'vacation/modals/modal_liquidation_list.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        employee = get_object_or_404(Employee.objects.select_related('person'), pk=self.kwargs['employee_id'])
        vacation_action_type = ActionType.objects.filter(name__iexact='VACACIONES').first()

        actions = PersonnelAction.objects.filter(
            employee=employee,
            action_type=vacation_action_type
        ).select_related('action_type', 'employee__person', 'vacation_request')

        search_query = self.request.GET.get('search', '').strip()
        if search_query:
            actions = actions.filter(number__icontains=search_query)

        actions = actions.order_by('-date_issue', '-id')
        paginator = Paginator(actions, 10)
        page_obj = paginator.page(self.request.GET.get('page', 1))

        context['employee'] = employee
        context['actions'] = page_obj
        context['search_query'] = search_query
        return context


class RegisterLiquidationView(LoginRequiredMixin, View):
    def post(self, request, action_id):
        try:
            action = get_object_or_404(PersonnelAction, pk=action_id)
            if action.is_registered:
                return JsonResponse({'status': 'error', 'success': False, 'message': 'Esta acción ya está registrada.'},
                                    status=400)

            vacation_request = VacationRequest.objects.get(personnel_action=action)

            with transaction.atomic():
                action.is_registered = True
                action.date_registered = dt.date.today()
                action.save()

                vacation_request.status = 'APPROVED'
                vacation_request.approved_by = request.user
                vacation_request.save()

                days_to_discount = float(vacation_request.days_quantity)
                VacationHistory.objects.create(
                    vacation_balance=vacation_request.balance_used,
                    vacation_request=vacation_request,
                    value_discount=days_to_discount,
                    days_discount=days_to_discount,
                    proportional_discount=0.0,
                    hours_discount=0.0,
                    minutes_discount=0.0,
                    observation=f'Liquidación de vacaciones registrada - Acción {action.number}',
                    created_by=request.user
                )

                EmployeeVacationBalance.objects.filter(id=vacation_request.balance_used.id).update(
                    balance_days=F('balance_days') - vacation_request.days_quantity,
                    vacation_days=F('vacation_days') + vacation_request.days_quantity
                )

            return JsonResponse(
                {'status': 'success', 'success': True, 'message': 'Liquidación registrada exitosamente.'})
        except Exception as e:
            return JsonResponse({'status': 'error', 'success': False, 'message': str(e)}, status=500)


class EditLiquidationView(LoginRequiredMixin, FormView):
    template_name = 'vacation/modals/modal_liquidation_edit.html'
    form_class = VacationLiquidationForm

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        action = get_object_or_404(PersonnelAction, pk=self.kwargs['action_id'])
        vacation_request = get_object_or_404(VacationRequest, personnel_action=action)
        context['action'] = action
        context['vacation_request'] = vacation_request
        context['employee'] = action.employee
        context['balance'] = vacation_request.balance_used
        context['action_url'] = reverse('vacation:edit_liquidation', kwargs={'action_id': action.id})
        return context

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        action = get_object_or_404(PersonnelAction, pk=self.kwargs['action_id'])
        vacation_request = get_object_or_404(VacationRequest, personnel_action=action)
        balance = vacation_request.balance_used
        kwargs['available_days'] = float(balance.balance_days + vacation_request.days_quantity) if balance else 0

        if not self.request.POST:
            kwargs['initial'] = {
                'start_date': vacation_request.start_date.strftime('%Y-%m-%d') if vacation_request.start_date else None,
                'end_date': vacation_request.end_date.strftime('%Y-%m-%d') if vacation_request.end_date else None,
                'nominating_authority': action.authority_1.id if action.authority_1 else None,
                'human_resources_responsible': action.authority_2.id if action.authority_2 else None,
                'registration_responsible': action.register.id if action.register else None,
                'review_responsible': action.reviewer.id if action.reviewer else None,
                'elaborated_by': action.elaboration.id if action.elaboration else None,
            }
        return kwargs

    def form_valid(self, form):
        try:
            action = get_object_or_404(PersonnelAction, pk=self.kwargs['action_id'])
            vacation_request = get_object_or_404(VacationRequest, personnel_action=action)

            if action.is_registered:
                return JsonResponse({'status': 'error', 'success': False,
                                     'message': 'No se puede editar una liquidación ya registrada.'}, status=400)

            start_date = form.cleaned_data['start_date']
            end_date = form.cleaned_data['end_date']
            days_requested = form.cleaned_data['days_requested']

            with transaction.atomic():
                vacation_request.start_date = start_date
                vacation_request.end_date = end_date
                vacation_request.days_quantity = Decimal(str(days_requested))
                vacation_request.save()

                action.date_effective = start_date
                action.authority_1 = form.cleaned_data['nominating_authority']
                action.authority_2 = form.cleaned_data['human_resources_responsible']
                action.register = form.cleaned_data['registration_responsible']
                action.reviewer = form.cleaned_data['review_responsible']
                action.elaboration = form.cleaned_data['elaborated_by']
                action.save()

            return JsonResponse(
                {'status': 'success', 'success': True, 'message': 'Liquidación actualizada exitosamente.'})
        except Exception as e:
            return JsonResponse({'status': 'error', 'success': False, 'message': str(e)}, status=500)

    def form_invalid(self, form):
        errors = {k: v[0] for k, v in form.errors.items()}
        return JsonResponse({
            'status': 'error',
            'success': False,
            'message': 'Por favor corrija los errores en el formulario',
            'errors': errors
        }, status=400)


# ==============================================================================
# HISTORIAL Y REPORTES
# ==============================================================================

class VacationHistoryDetailView(LoginRequiredMixin, TemplateView):
    template_name = 'vacation/modals/modal_vacation_history.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        balance = get_object_or_404(
            EmployeeVacationBalance.objects.select_related('employee__person', 'period'),
            pk=self.kwargs.get('balance_id')
        )
        history_records = VacationHistory.objects.filter(
            vacation_balance=balance,
            vacation_request__isnull=False
        ).select_related('vacation_request', 'vacation_request__personnel_action').order_by('-created_at')

        context['employee'] = balance.employee
        context['period'] = balance.period
        context['balance'] = balance
        context['history_records'] = history_records
        return context


class PermitHistoryDetailView(LoginRequiredMixin, TemplateView):
    template_name = 'vacation/modals/modal_permit_history.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        balance = get_object_or_404(
            EmployeeVacationBalance.objects.select_related('employee__person', 'period'),
            pk=self.kwargs.get('balance_id')
        )
        history_records = VacationHistory.objects.filter(
            vacation_balance=balance,
            permit_request__isnull=False
        ).select_related('permit_request').order_by('-created_at')

        context['employee'] = balance.employee
        context['period'] = balance.period
        context['balance'] = balance
        context['history_records'] = history_records
        return context


class PermitReportModalView(LoginRequiredMixin, TemplateView):
    template_name = 'vacation/modals/modal_permit_report.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        employee = get_object_or_404(Employee.objects.select_related('person'), pk=self.kwargs.get('employee_id'))
        context['employee'] = employee
        return context


class PermitReportPDFView(LoginRequiredMixin, View):
    def get(self, request, employee_id):
        start_date_str = request.GET.get('start_date')
        end_date_str = request.GET.get('end_date')

        if not start_date_str or not end_date_str:
            return HttpResponse('Fechas no proporcionadas', status=400)

        start_date = dt.datetime.strptime(start_date_str, '%Y-%m-%d').date()
        end_date = dt.datetime.strptime(end_date_str, '%Y-%m-%d').date()
        employee = get_object_or_404(Employee.objects.select_related('person', 'area'), pk=employee_id)

        permits = VacationHistory.objects.filter(
            vacation_balance__employee=employee,
            permit_request__isnull=False,
            permit_request__start_date__gte=start_date,
            permit_request__start_date__lte=end_date,
            permit_request__status='APPROVED'
        ).select_related(
            'permit_request',
            'permit_request__permit_type',
            'vacation_balance__period'
        ).order_by('permit_request__start_date')

        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet

        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter)
        elements = []
        styles = getSampleStyleSheet()

        elements.append(Paragraph('REPORTE DE PERMISOS CON CARGO A VACACIONES', styles['Heading1']))
        elements.append(Paragraph(f'Empleado: {employee.person.full_name}', styles['Normal']))
        elements.append(Paragraph(f'Período: {start_date} al {end_date}', styles['Normal']))
        elements.append(Spacer(1, 20))

        doc.build(elements)
        buffer.seek(0)
        response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = f'inline; filename="Reporte_Permisos_{employee_id}.pdf"'
        return response


class LiquidationPrintPDFView(LoginRequiredMixin, View):
    def get(self, request, action_id):
        from xhtml2pdf import pisa
        action = get_object_or_404(
            PersonnelAction.objects.select_related('employee__person', 'action_type'),
            pk=action_id
        )
        vacation_request = getattr(action, 'vacation_request', None)

        if not vacation_request:
            return HttpResponse('Esta acción no está relacionada con liquidación de vacaciones', status=400)

        template = get_template('vacation/reports/pdf_liquidation.html')
        html = template.render({
            'action': action,
            'employee': action.employee,
            'vacation_request': vacation_request,
            'today': dt.datetime.now()
        })

        buffer = BytesIO()
        pdf = pisa.pisaDocument(BytesIO(html.encode("UTF-8")), buffer, encoding='UTF-8')
        if not pdf.err:
            response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
            response['Content-Disposition'] = f'inline; filename="Liquidacion_{action.number}.pdf"'
            return response
        return HttpResponse('Error al generar PDF', status=500)


# # ==============================================================================
# # CRONOGRAMA DE VACACIONES
# # ==============================================================================

class VacationSchedulePlanListView(LoginRequiredMixin, PermissionRequiredMixin, ListView):
    model = VacationSchedulePlan
    template_name = 'vacation/schedule_plan_list.html'
    permission_required = 'vacation.view_vacationscheduleplan'
    context_object_name = 'plans'
    paginate_by = 10

    def get_queryset(self):
        show_inactive = self.request.GET.get('show_inactive') == 'true'

        if show_inactive:
            qs = VacationSchedulePlan.objects.filter(status=VacationSchedulePlan.PlanStatus.CLOSED)
        else:
            qs = VacationSchedulePlan.objects.exclude(status=VacationSchedulePlan.PlanStatus.CLOSED)

        q = self.request.GET.get('q', '').strip()
        if q:
            qs = qs.filter(Q(title__icontains=q) | Q(year__icontains=q))

        # ANOTACIÓN: Conteo de solicitudes pendientes para la campana en el partial
        qs = qs.annotate(
            pending_discrepancies_count=Count(
                'discrepancy_requests',
                filter=Q(discrepancy_requests__status='PENDING')
            )
        )

        return qs.order_by('-year')


class VacationSchedulePlanCreateModalView(LoginRequiredMixin, CreateView, PermissionRequiredMixin):
    """Modal para aperturar cronograma anual."""
    model = VacationSchedulePlan
    form_class = VacationSchedulePlanForm
    template_name = 'vacation/modals/modal_schedule_plan_form.html'
    success_url = reverse_lazy('vacation:schedule_plan_list')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['modal_title'] = 'Aperturar Plan Anual'
        context['action_url'] = reverse('vacation:schedule_plan_create')
        return context

    def form_valid(self, form):
        form.instance.created_by = self.request.user
        self.object = form.save()
        if self.request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return JsonResponse({
                'status': 'success',
                'success': True,
                'message': f'Cronograma {self.object.year} creado con éxito.',
                'redirect_url': str(self.success_url)
            })
        messages.success(self.request, f'Cronograma {self.object.year} creado con éxito.')
        return redirect(self.success_url)

    def form_invalid(self, form):
        if self.request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            errors = {k: v[0] for k, v in form.errors.items()}
            return JsonResponse({'status': 'error', 'success': False, 'errors': errors}, status=400)
        return super().form_invalid(form)


class VacationSchedulePlanUpdateModalView(LoginRequiredMixin, UpdateView, PermissionRequiredMixin):
    """Modal para editar cronograma anual (reutiliza el mismo template)."""
    model = VacationSchedulePlan
    form_class = VacationSchedulePlanForm
    template_name = 'vacation/modals/modal_schedule_plan_form.html'
    success_url = reverse_lazy('vacation:schedule_plan_list')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['modal_title'] = f'Editar Plan Anual ({self.object.year})'
        context['action_url'] = reverse('vacation:schedule_plan_edit', kwargs={'pk': self.object.pk})
        return context

    def dispatch(self, request, *args, **kwargs):
        obj = self.get_object()
        if obj.status == VacationSchedulePlan.PlanStatus.APPROVED:
            return JsonResponse({
                'status': 'error',
                'success': False,
                'message': 'No se puede editar un cronograma que ya ha sido Aprobado Oficialmente.'
            }, status=403)
        return super().dispatch(request, *args, **kwargs)

    def form_valid(self, form):
        self.object = form.save()
        if self.request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return JsonResponse({
                'status': 'success',
                'success': True,
                'message': f'Cronograma {self.object.year} actualizado con éxito.',
                'redirect_url': str(self.success_url)
            })
        messages.success(self.request, f'Cronograma {self.object.year} actualizado con éxito.')
        return redirect(self.success_url)

    def form_invalid(self, form):
        if self.request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            errors = {k: v[0] for k, v in form.errors.items()}
            return JsonResponse({'status': 'error', 'success': False, 'errors': errors}, status=400)
        return super().form_invalid(form)


class VacationScheduleDetailView(LoginRequiredMixin, TemplateView, PermissionRequiredMixin):
    template_name = 'vacation/schedule_detail.html'
    permission_required = 'vacation.view_vacationscheduleplan'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        plan = get_object_or_404(VacationSchedulePlan, pk=self.kwargs['plan_id'])
        user = self.request.user

        # 1. Resolver unidad principal ascendente y todos sus descendientes
        base_unit, unit_ids = resolve_scheduling_hierarchy(
            user=user,
            requested_unit_id=self.request.GET.get('unit_id')
        )

        # 2. Filtrar todos los empleados que pertenecen a la unidad principal o sus hijas
        employees_qs = Employee.objects.filter(
            is_active=True,
            area_id__in=unit_ids
        ).exclude(
            Q(employment_status__code__icontains='PROFESIONAL') |
            Q(employment_status__name__icontains='PROFESIONAL') |
            Q(employment_status__code__icontains='SERVICIOS') |
            Q(employment_status__name__icontains='SERVICIOS') |
            Q(employment_status__code__icontains='HONORARIOS') |
            Q(employment_status__name__icontains='HONORARIOS')
        ).select_related(
            'person',
            'area',
            'employment_status',
            'institutional_data'
        ).order_by('person__last_name', 'person__first_name')

        scheduled_items = VacationScheduleItem.objects.filter(
            plan=plan,
            employee__in=employees_qs
        ).select_related('employee__person')

        # 3. Meses del año
        months_cards = []
        for m in range(1, 13):
            items_in_m = [it for it in scheduled_items if it.start_date.month == m]
            months_cards.append({
                'number': m,
                'name': MONTH_NAMES.get(m, ''),
                'items': items_in_m
            })

        # 4. Obtener IDs de empleados reportados como "No pertenece"
        reported_emp_ids = set(
            ScheduleDiscrepancyRequest.objects.filter(
                plan=plan,
                request_type=ScheduleDiscrepancyRequest.RequestType.WRONG_UNIT,
                status=ScheduleDiscrepancyRequest.Status.PENDING
            ).values_list('employee_id', flat=True)
        )

        # 4.1 Obtener empleados SOLICITADOS pendientes de incorporar a esta unidad
        requested_discrepancies = ScheduleDiscrepancyRequest.objects.filter(
            plan=plan,
            unit=base_unit,
            request_type=ScheduleDiscrepancyRequest.RequestType.MISSING_EMPLOYEE,
            status=ScheduleDiscrepancyRequest.Status.PENDING
        ).select_related('employee__person', 'employee__employment_status')

        # 5. Banquillo con cálculo proporcional y banderas
        employee_cards = []
        for emp in employees_qs:
            allowed_days, ceiling, is_proportional, display_val = calculate_available_days_for_month(
                emp, plan.year, 1
            )
            total_planned = sum(it.days_quantity for it in scheduled_items if it.employee_id == emp.id)
            max_assignable_plan = ceiling if is_proportional else allowed_days
            display_max = f"{display_val}+" if is_proportional else f"{allowed_days}"

            if total_planned < max_assignable_plan:
                remaining = max_assignable_plan - total_planned
                employee_cards.append({
                    'employee': emp,
                    'total_planned_days': total_planned,
                    'remaining_days': remaining,
                    'max_days': max_assignable_plan,
                    'display_max': display_max,
                    'is_proportional': is_proportional,
                    'is_reported_wrong': emp.id in reported_emp_ids,
                    'is_requested_ghost': False
                })

        # 5.1 Agregar las tarjetas "fantasma" solicitadas
        for disc in requested_discrepancies:
            emp = disc.employee
            # Si no está ya en la lista
            if not any(c['employee'].id == emp.id for c in employee_cards):
                allowed_days, ceiling, is_proportional, display_val = calculate_available_days_for_month(
                    emp, plan.year, 1
                )
                employee_cards.append({
                    'employee': emp,
                    'total_planned_days': 0,
                    'remaining_days': allowed_days,
                    'max_days': allowed_days,
                    'display_max': display_val,
                    'is_proportional': is_proportional,
                    'is_reported_wrong': False,
                    'is_requested_ghost': True,
                })

        context['plan'] = plan
        context['base_unit'] = base_unit
        context['employee_cards'] = employee_cards
        context['months_cards'] = months_cards
        return context


class VacationScheduleDropModalView(LoginRequiredMixin, PermissionRequiredMixin, View):
    template_name = 'vacation/modals/modal_schedule_drop_form.html'
    permission_required = 'vacation.view_vacationscheduleplan'

    def get(self, request, plan_id, employee_id, month):
        plan = get_object_or_404(VacationSchedulePlan, pk=plan_id)
        employee = get_object_or_404(
            Employee.objects.select_related('person', 'employment_status'),
            pk=employee_id
        )

        month_int = int(month)

        # Usar la función probada en el shell
        allowed_int, ceiling, is_proportional, display_val = calculate_available_days_for_month(
            employee, plan.year, month_int
        )
        annual_cap = allowed_int

        # Sumar los días ya planificados en este cronograma
        planned_so_far = VacationScheduleItem.objects.filter(
            plan=plan,
            employee=employee
        ).aggregate(total=Sum('days_quantity'))['total'] or 0

        # Cupo restante exacto para este mes
        max_assignable = max(0, annual_cap - planned_so_far)
        default_days = min(15, max_assignable) if max_assignable > 0 else 0

        context = {
            'plan': plan,
            'employee': employee,
            'month': month_int,
            'month_name': MONTH_NAMES.get(month_int, ''),
            'annual_cap': annual_cap,
            'planned_so_far': planned_so_far,
            'max_assignable': max_assignable,
            'default_days': default_days,
            'action_url': reverse('vacation:schedule_drop_save',
                                  kwargs={'plan_id': plan.id, 'employee_id': employee.id})
        }
        return render(request, self.template_name, context)


# vacation/views.py

class VacationScheduleDropSaveView(LoginRequiredMixin, PermissionRequiredMixin, View):
    permission_required = 'vacation.view_vacationscheduleplan'

    def post(self, request, plan_id, employee_id):
        plan = get_object_or_404(VacationSchedulePlan, pk=plan_id)
        employee = get_object_or_404(Employee.objects.select_related('employment_status'), pk=employee_id)

        if plan.status == VacationSchedulePlan.PlanStatus.APPROVED:
            return JsonResponse({
                'status': 'error',
                'success': False,
                'message': 'El cronograma está Aprobado Oficialmente y no admite modificaciones.'
            }, status=400)

        month_str = request.POST.get('month')
        days_str = request.POST.get('days_quantity') or request.POST.get('days')
        observation = request.POST.get('observation', '').strip()

        try:
            month = int(month_str)
            days_quantity = int(days_str)
        except (ValueError, TypeError):
            return JsonResponse({'status': 'error', 'success': False, 'message': 'Datos numéricos inválidos.'},
                                status=400)

        if days_quantity < 1:
            return JsonResponse({'status': 'error', 'success': False, 'message': 'Debe asignar al menos 1 día.'},
                                status=400)

        # Validación con la misma función de services.py
        allowed_int, ceiling, is_proportional, _ = calculate_available_days_for_month(
            employee, plan.year, month
        )
        annual_cap = allowed_int

        # Evitar asignación duplicada en el mismo mes
        already_in_month = VacationScheduleItem.objects.filter(
            plan=plan,
            employee=employee,
            start_date__month=month
        ).exists()

        if already_in_month:
            return JsonResponse({
                'status': 'error',
                'success': False,
                'message': 'El funcionario ya cuenta con vacaciones planificadas en este mes.'
            }, status=400)

        planned_so_far = VacationScheduleItem.objects.filter(
            plan=plan,
            employee=employee
        ).aggregate(total=Sum('days_quantity'))['total'] or 0

        disponible = max(0, annual_cap - planned_so_far)

        if days_quantity > disponible:
            return JsonResponse({
                'status': 'error',
                'success': False,
                'message': f'Excede el límite. Ya programados: {planned_so_far}/{annual_cap}. Días disponibles: {disponible}.'
            }, status=400)

        max_day = calendar.monthrange(plan.year, month)[1]
        start_date = dt.date(plan.year, month, 1)
        end_date = dt.date(plan.year, month, min(days_quantity, max_day))

        VacationScheduleItem.objects.create(
            plan=plan,
            employee=employee,
            start_date=start_date,
            end_date=end_date,
            days_quantity=days_quantity,
            status=VacationScheduleItem.ItemStatus.PROPOSED,
            observation=observation or f'{days_quantity} días programados para {MONTH_NAMES.get(month, "")}',
            created_by=request.user
        )

        return JsonResponse({
            'status': 'success',
            'success': True,
            'message': f'Se programaron {days_quantity} días exitosamente.'
        })


class VacationScheduleItemCreateModalView(LoginRequiredMixin, View):
    """Modal para asignar/planificar fechas en el cronograma."""
    template_name = 'vacation/modals/modal_schedule_item_form.html'

    def get(self, request, plan_id):
        plan = get_object_or_404(VacationSchedulePlan, pk=plan_id)
        form = VacationScheduleItemForm()
        context = {
            'plan': plan,
            'form': form,
            'action_url': reverse('vacation:schedule_item_create', kwargs={'plan_id': plan.id})
        }
        # Si es petición AJAX desde openAjaxModal
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            html = render_to_string(self.template_name, context, request=request)
            return HttpResponse(html)
        return render(request, self.template_name, context)

    def post(self, request, plan_id):
        plan = get_object_or_404(VacationSchedulePlan, pk=plan_id)
        form = VacationScheduleItemForm(request.POST)

        if form.is_valid():
            item = form.save(commit=False)
            item.plan = plan
            item.days_quantity = form.cleaned_data['days_quantity']
            item.created_by = request.user
            item.save()

            return JsonResponse({
                'status': 'success',
                'success': True,
                'message': f'Fracción de {item.days_quantity} días programada con éxito.',
                'redirect_url': reverse('vacation:schedule_detail', kwargs={'plan_id': plan.id})
            })

        errors = {k: v[0] for k, v in form.errors.items()}
        return JsonResponse({
            'status': 'error',
            'success': False,
            'message': 'Por favor revise los errores del formulario.',
            'errors': errors
        }, status=400)


class VacationScheduleItemDeleteView(LoginRequiredMixin, View):
    """Eliminar ítem del cronograma vía AJAX."""

    def post(self, request, item_id):
        item = get_object_or_404(VacationScheduleItem, pk=item_id)
        item.delete()
        return JsonResponse({
            'status': 'success',
            'success': True,
            'message': 'Fracción de vacación removida del cronograma.'
        })


class VacationSchedulePlanToggleStatusView(LoginRequiredMixin, PermissionRequiredMixin, View):
    """Cierra o abre el cronograma anual de vacaciones."""

    def post(self, request, pk):
        plan = get_object_or_404(VacationSchedulePlan, pk=pk)
        if plan.status == VacationSchedulePlan.PlanStatus.APPROVED:
            return JsonResponse({
                'status': 'error',
                'success': False,
                'message': 'Un cronograma Aprobado Oficialmente no puede ser dado de baja ni modificado.'
            }, status=400)

        if plan.status == VacationSchedulePlan.PlanStatus.CLOSED:
            plan.status = VacationSchedulePlan.PlanStatus.OPEN
            action_word = "reaperturado"
            is_active = True
        else:
            plan.status = VacationSchedulePlan.PlanStatus.CLOSED
            action_word = "dado de baja / cerrado"
            is_active = False
        plan.save()

        return JsonResponse({
            'status': 'success',
            'success': True,
            'message': f'El cronograma {plan.year} ha sido {action_word} exitosamente.',
            'is_active': is_active
        })


class VacationScheduleExcelExportView(LoginRequiredMixin, View):
    """
    Genera el reporte en Excel (.xlsx) con estilos corporativos del cronograma
    respetando la dependencia principal cabecera y sus firmas institucionales.
    """

    def get(self, request, plan_id):
        plan = get_object_or_404(VacationSchedulePlan, pk=plan_id)

        # 1. Resolver unidad principal cabecera (is_main_unit) y todas sus subdependencias
        base_unit, unit_ids = resolve_scheduling_hierarchy(
            user=request.user,
            requested_unit_id=request.GET.get('unit_id')
        )

        # 2. Filtrar empleados de la unidad principal y sus hijas (excluyendo servicios profesionales)
        active_contract_prefetch = Prefetch(
            'management_periods',
            queryset=ManagementPeriod.objects.filter(
                status__code='ACTIVO'
            ).select_related('contract_type__labor_regime'),
            to_attr='active_contracts'
        )

        employees_qs = Employee.objects.filter(is_active=True)
        if unit_ids:
            employees_qs = employees_qs.filter(area_id__in=unit_ids)

        employees = employees_qs.exclude(
            employment_status__code__icontains='SERVICIOS_PROFESIONALES'
        ).exclude(
            employment_status__name__icontains='SERVICIOS PROFESIONALES'
        ).select_related(
            'person',
            'employment_status',
            'area',
            'institutional_data'
        ).prefetch_related(
            active_contract_prefetch
        ).order_by('person__last_name', 'person__first_name')

        scheduled_items = VacationScheduleItem.objects.filter(
            plan=plan,
            employee__in=employees
        ).select_related('employee__person')

        month_names = {
            1: 'Enero', 2: 'Febrero', 3: 'Marzo', 4: 'Abril',
            5: 'Mayo', 6: 'Junio', 7: 'Julio', 8: 'Agosto',
            9: 'Septiembre', 10: 'Octubre', 11: 'Noviembre', 12: 'Diciembre'
        }

        # 3. Construir Excel
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Cronograma {plan.year}"
        ws.views.sheetView[0].showGridLines = True

        # Paleta de estilos corporativos
        header_fill = PatternFill(start_color="203C7D", end_color="203C7D", fill_type="solid")
        header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        title_font = Font(name="Arial", size=12, bold=True, color="1E293B")
        sub_font = Font(name="Arial", size=9, bold=False, color="64748B")
        regular_font = Font(name="Arial", size=9)
        bold_font = Font(name="Arial", size=9, bold=True)
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        # Encabezado Institucional
        ws.merge_cells("A1:G1")
        ws["A1"] = f"CRONOGRAMA ANUAL DE VACACIONES {plan.year} - {plan.title.upper()}"
        ws["A1"].font = title_font
        ws["A1"].alignment = left_align

        ws.merge_cells("A2:G2")
        ws[
            "A2"] = f"Dependencia Principal: {base_unit.name if base_unit else 'Todas las dependencias'} | Fecha de emisión: {dt.date.today().strftime('%d/%m/%Y')}"
        ws["A2"].font = sub_font
        ws["A2"].alignment = left_align

        ws.row_dimensions[1].height = 24
        ws.row_dimensions[2].height = 18
        ws.row_dimensions[4].height = 26

        # Cabecera de la tabla
        headers = [
            "N°", "CÉDULA", "APELLIDOS Y NOMBRES", "FECHA INGRESO",
            "RÉGIMEN / TIPO", "MESES PROGRAMADOS", "TOTAL DÍAS"
        ]

        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=4, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        # Llenar registros
        current_row = 5
        secuencial = 1

        for emp in employees:
            items = [it for it in scheduled_items if it.employee_id == emp.id]
            if not items:
                continue

            inst_data = getattr(emp, 'institutional_data', None)
            fecha_ingreso = inst_data.entry_date.strftime('%d/%m/%Y') if (inst_data and inst_data.entry_date) else '-'

            # Apellidos y Nombres en orden correcto
            apellidos_nombres = f"{emp.person.last_name} {emp.person.first_name}".strip()

            # Régimen / Tipo
            active_contract = emp.active_contracts[0] if getattr(emp, 'active_contracts', None) else None
            regimen_nombre = "LOSEP"
            if active_contract and active_contract.contract_type and active_contract.contract_type.labor_regime:
                regimen_nombre = active_contract.contract_type.labor_regime.name or active_contract.contract_type.labor_regime.code

            status_text = (emp.employment_status.name or '') if emp.employment_status else ''
            tipo_display = "TRABAJADOR" if any(
                k in status_text.upper() for k in ['TRABAJADOR', 'OBRERO', 'CODIGO']) else "EMPLEADO"
            regimen_tipo = f"{regimen_nombre} / {tipo_display}"

            # Meses programados
            meses_str = ", ".join([f"{month_names.get(it.start_date.month, '')} ({it.days_quantity}d)" for it in items])
            total_dias = sum(it.days_quantity for it in items)

            row_data = [
                secuencial,
                emp.person.document_number,
                apellidos_nombres,
                fecha_ingreso,
                regimen_tipo,
                meses_str,
                f"{total_dias} días"
            ]

            for col_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.font = regular_font
                cell.border = thin_border

                if col_idx in [1, 2, 4, 7]:
                    cell.alignment = center_align
                    if col_idx == 7:
                        cell.font = bold_font
                else:
                    cell.alignment = left_align

            ws.row_dimensions[current_row].height = 20
            current_row += 1
            secuencial += 1

        # Bloque de Firmas Institucionales (Elaborado por y Autorizado por)
        current_row += 2
        sig_header_row = current_row
        ws.cell(row=sig_header_row, column=2, value="ELABORADO POR:").font = bold_font
        ws.cell(row=sig_header_row, column=5, value="AUTORIZADO POR:").font = bold_font
        ws.cell(row=sig_header_row, column=2).alignment = center_align
        ws.cell(row=sig_header_row, column=5).alignment = center_align

        # Espacio para firmas y nombres
        sig_name_row = current_row + 4
        elaborado_nombre = request.user.signature_name
        elaborado_cargo = request.user.signature_position or "Funcionario Responsable"

        autorizado_nombre = "SIN ASIGNAR"
        autorizado_cargo = "Responsable de Dependencia"
        if base_unit and base_unit.boss and base_unit.boss.person:
            boss_user = getattr(base_unit.boss.person, 'user', None)
            if boss_user:
                autorizado_nombre = boss_user.signature_name
                autorizado_cargo = boss_user.signature_position or f"Jefe / Responsable de {base_unit.name}"
            else:
                autorizado_nombre = base_unit.boss.person.full_name.upper()
                autorizado_cargo = f"Responsable de {base_unit.name}"

        # Nombres y Cargos de los firmantes
        ws.cell(row=sig_name_row, column=2, value=elaborado_nombre).font = bold_font
        ws.cell(row=sig_name_row + 1, column=2, value=elaborado_cargo).font = regular_font
        ws.cell(row=sig_name_row, column=2).alignment = center_align
        ws.cell(row=sig_name_row + 1, column=2).alignment = center_align

        ws.cell(row=sig_name_row, column=5, value=autorizado_nombre).font = bold_font
        ws.cell(row=sig_name_row + 1, column=5, value=autorizado_cargo).font = regular_font
        ws.cell(row=sig_name_row, column=5).alignment = center_align
        ws.cell(row=sig_name_row + 1, column=5).alignment = center_align

        # Anchos de columna proporcionales
        column_widths = [6, 15, 38, 16, 26, 36, 14]
        for idx, width in enumerate(column_widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width

        # Descarga
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        unit_filename = base_unit.name.replace(' ', '_') if base_unit else 'General'
        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response[
            "Content-Disposition"] = f'attachment; filename="Cronograma_Vacaciones_{unit_filename}_{plan.year}.xlsx"'
        return response


# vacation/views.py

class VacationSchedulePdfExportView(LoginRequiredMixin, View):
    """
    Genera el reporte en PDF vertical con la dependencia principal resuelta y su jefe firmante.
    """

    def get(self, request, plan_id):
        plan = get_object_or_404(VacationSchedulePlan, pk=plan_id)

        # 1. Resolver unidad principal
        base_unit, unit_ids = resolve_scheduling_hierarchy(
            user=request.user,
            requested_unit_id=request.GET.get('unit_id')
        )

        # 2. Empleados de toda la unidad principal y sus hijas
        employees_qs = Employee.objects.filter(
            is_active=True,
            area_id__in=unit_ids
        ).exclude(
            employment_status__code__icontains='SERVICIOS_PROFESIONALES'
        ).exclude(
            employment_status__name__icontains='SERVICIOS PROFESIONALES'
        ).select_related(
            'person',
            'employment_status',
            'institutional_data'
        ).order_by('person__last_name', 'person__first_name')

        scheduled_items = VacationScheduleItem.objects.filter(
            plan=plan,
            employee__in=employees_qs
        ).select_related('employee__person')

        # 3. Filas del reporte
        report_rows = []
        secuencial = 1
        for emp in employees_qs:
            items = [it for it in scheduled_items if it.employee_id == emp.id]
            if not items:
                continue

            meses_str = ", ".join([f"{MONTH_NAMES.get(it.start_date.month, '')} ({it.days_quantity}d)" for it in items])
            total_dias = sum(it.days_quantity for it in items)
            apellidos_nombres = f"{emp.person.last_name} {emp.person.first_name}".strip()

            tipo_display = (
                (emp.employment_status.name or emp.employment_status.code or '').strip().upper()
                if emp.employment_status else 'SIN TIPO'
            )

            inst_data = getattr(emp, 'institutional_data', None)
            fecha_ingreso = inst_data.entry_date.strftime('%d/%m/%Y') if (inst_data and inst_data.entry_date) else '-'

            report_rows.append({
                'secuencial': secuencial,
                'cedula': emp.person.document_number,
                'apellidos_nombres': apellidos_nombres,
                'fecha_ingreso': fecha_ingreso,
                'tipo': tipo_display,
                'meses': meses_str,
                'total_dias': total_dias,
            })
            secuencial += 1

        sys_config = SystemConfiguration.get_current()
        letterhead_path = sys_config.letterhead.path if sys_config and sys_config.letterhead else None

        # 4. Autoridades Firmantes
        elaborado_nombre = request.user.signature_name
        elaborado_cargo = request.user.signature_position or "Funcionario Responsable"

        # AUTORIZADO POR: Siempre el jefe de la base_unit resuelta (unidad principal)
        autorizado_nombre = "SIN ASIGNAR"
        autorizado_cargo = "Responsable de Dependencia"
        if base_unit and base_unit.boss and base_unit.boss.person:
            boss_user = getattr(base_unit.boss.person, 'user', None)
            if boss_user:
                autorizado_nombre = boss_user.signature_name
                autorizado_cargo = boss_user.signature_position or f"Jefe / Responsable de {base_unit.name}"
            else:
                autorizado_nombre = base_unit.boss.person.full_name.upper()
                autorizado_cargo = f"Responsable de {base_unit.name}"

        template = get_template('vacation/reports/pdf_schedule_report.html')
        html = template.render({
            'plan': plan,
            'base_unit': base_unit,
            'sys_config': sys_config,
            'letterhead_path': letterhead_path,
            'report_rows': report_rows,
            'today': dt.date.today(),
            'total_funcionarios': len(report_rows),
            'elaborado_nombre': elaborado_nombre,
            'elaborado_cargo': elaborado_cargo,
            'autorizado_nombre': autorizado_nombre,
            'autorizado_cargo': autorizado_cargo,
        })

        buffer = BytesIO()
        pdf = pisa.pisaDocument(BytesIO(html.encode("UTF-8")), buffer, encoding='UTF-8')
        if not pdf.err:
            response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
            response[
                'Content-Disposition'] = f'inline; filename="Cronograma_Vacaciones_{base_unit.name}_{plan.year}.pdf"'
            return response

        return HttpResponse('Error al generar el reporte PDF', status=500)


class VacationScheduleConsolidatedModalView(LoginRequiredMixin, View, PermissionRequiredMixin):
    """
    Despliega el modal AJAX para configurar el reporte consolidado.
    """
    template_name = 'vacation/modals/modal_consolidated_report.html'

    def get(self, request, plan_id):
        plan = get_object_or_404(VacationSchedulePlan, pk=plan_id)
        sys_config = SystemConfiguration.get_current()
        current_u = request.user
        current_u_text = f"{current_u.signature_name.upper()}" + (
            f" - {current_u.signature_position}" if current_u.signature_position else "")

        context = {
            'plan': plan,
            'sys_config': sys_config,
            'current_user_id': current_u.id,
            'current_user_text': current_u_text,
            'action_url': reverse('vacation:schedule_consolidated_export', kwargs={'plan_id': plan.id})
        }
        return render(request, self.template_name, context)


class VacationScheduleConsolidatedExportView(LoginRequiredMixin, View, PermissionRequiredMixin):
    """
    Genera el reporte consolidado institucional en PDF o Excel según lo solicitado.
    """

    def post(self, request, plan_id):
        plan = get_object_or_404(VacationSchedulePlan, pk=plan_id)
        form = ConsolidatedReportForm(request.POST)

        if not form.is_valid():
            return JsonResponse({'status': 'error', 'message': 'Datos inválidos en el formulario.'}, status=400)

        report_format = form.cleaned_data['report_format']
        elaborated_user = form.cleaned_data['elaborated_by']
        reviewed_user = form.cleaned_data['reviewed_by']

        # 1. Traer TODOS los empleados de la institución que no sean SERVICIOS PROFESIONALES
        active_contract_prefetch = Prefetch(
            'management_periods',
            queryset=ManagementPeriod.objects.filter(
                Q(status__code__icontains='FIRMADO') |
                Q(status__name__icontains='FIRMADO')
            ).select_related(
                'contract_type__labor_regime'
            ).order_by('-start_date', '-pk'),
            to_attr='active_contracts'
        )

        employees = Employee.objects.filter(is_active=True).exclude(
            employment_status__code__icontains='SERVICIOS_PROFESIONALES'
        ).exclude(
            employment_status__name__icontains='SERVICIOS PROFESIONALES'
        ).select_related(
            'person',
            'employment_status',
            'area',
            'institutional_data'
        ).prefetch_related(
            active_contract_prefetch
        ).order_by('person__last_name', 'person__first_name')

        # 2. Asignaciones de vacaciones en este plan
        scheduled_items = VacationScheduleItem.objects.filter(
            plan=plan,
            employee__in=employees
        ).select_related('employee')

        full_month_names = {
            1: 'Enero', 2: 'Febrero', 3: 'Marzo', 4: 'Abril',
            5: 'Mayo', 6: 'Junio', 7: 'Julio', 8: 'Agosto',
            9: 'Septiembre', 10: 'Octubre', 11: 'Noviembre', 12: 'Diciembre'
        }

        # 3. Estructurar filas consolidadas
        report_rows = []
        secuencial = 1
        for emp in employees:
            items = [it for it in scheduled_items if it.employee_id == emp.id]
            if not items:
                continue

            meses_str = ", ".join(
                [f"{full_month_names.get(it.start_date.month, '')} ({it.days_quantity}d)" for it in items]
            )
            total_dias = sum(it.days_quantity for it in items)
            apellidos_nombres = f"{emp.person.last_name} {emp.person.first_name}".strip()

            # Datos laborales reales del empleado.
            tipo_limpio = (
                (emp.employment_status.name or emp.employment_status.code or '').strip().upper()
                if emp.employment_status else 'SIN TIPO'
            )
            _, dias_ley, _ = get_vacation_regime_info(emp)
            inst_data = getattr(emp, 'institutional_data', None)
            fecha_ingreso = (
                inst_data.entry_date.strftime('%d/%m/%Y')
                if inst_data and inst_data.entry_date else '-'
            )
            active_contract = emp.active_contracts[0] if emp.active_contracts else None
            regimen = 'SIN RÉGIMEN'
            if active_contract and active_contract.contract_type:
                labor_regime = active_contract.contract_type.labor_regime
                if labor_regime:
                    regimen = (
                            labor_regime.name or labor_regime.code or regimen
                    ).strip().upper()
            unidad = (emp.area.name if emp.area else 'SIN ASIGNAR').strip().upper()

            report_rows.append({
                'secuencial': secuencial,
                'cedula': emp.person.document_number,
                'apellidos_nombres': apellidos_nombres,
                'regimen': regimen,
                'fecha_ingreso': fecha_ingreso,
                'unidad': unidad,
                'tipo': tipo_limpio,
                'dias_ley': f'{dias_ley} días',
                'meses': meses_str,
                'total_dias': total_dias,
            })
            secuencial += 1

        # 4. Configuración del sistema y 4 Firmas
        sys_config = SystemConfiguration.get_current()

        firmas = {
            'elaborado_nombre': elaborated_user.signature_name,
            'elaborado_cargo': elaborated_user.signature_position or "SERVIDOR DE TALENTO HUMANO",
            'revisado_nombre': reviewed_user.signature_name,
            'revisado_cargo': reviewed_user.signature_position or "COORDINADOR DE BASE DE DATOS Y NOMINA",
            'visto_bueno_nombre': sys_config.talento_humano_authority_name if sys_config else "DIRECTOR DE TALENTO HUMANO",
            'visto_bueno_cargo': sys_config.talento_humano_authority_position if sys_config else "DIRECTOR DE TALENTO HUMANO",
            'aprobado_nombre': sys_config.max_authority_name if sys_config else "MAXIMA AUTORIDAD",
            'aprobado_cargo': sys_config.max_authority_position if sys_config else "ALCALDE/SA",
        }

        # 5. Exportar a EXCEL
        if report_format == 'excel':
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = f"Consolidado {plan.year}"
            ws.views.sheetView[0].showGridLines = True

            header_fill = PatternFill(start_color="203C7D", end_color="203C7D", fill_type="solid")
            header_font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
            title_font = Font(name="Arial", size=12, bold=True, color="1E293B")
            sub_font = Font(name="Arial", size=8.5, bold=False, color="64748B")
            regular_font = Font(name="Arial", size=8.5)
            bold_font = Font(name="Arial", size=8.5, bold=True)
            center_align = Alignment(horizontal="center", vertical="center")
            left_align = Alignment(horizontal="left", vertical="center")
            thin_border = Border(
                left=Side(style='thin', color='CBD5E1'),
                right=Side(style='thin', color='CBD5E1'),
                top=Side(style='thin', color='CBD5E1'),
                bottom=Side(style='thin', color='CBD5E1')
            )

            # Título institucional
            ws.merge_cells("A1:I1")
            ws["A1"] = f"CONSOLIDADO INSTITUCIONAL DE VACACIONES {plan.year} - {plan.title.upper()}"
            ws["A1"].font = title_font
            ws["A1"].alignment = left_align

            ws.merge_cells("A2:I2")
            ws[
                "A2"] = f"Ámbito: Todas las Dependencias | Fecha de Emisión: {dt.date.today().strftime('%d/%m/%Y')} | Total Servidores: {len(report_rows)}"
            ws["A2"].font = sub_font
            ws["A2"].alignment = left_align

            headers = [
                "N°", "CÉDULA", "APELLIDOS Y NOMBRES", "RÉGIMEN",
                "FECHA INGRESO", "UNIDAD ADMINISTRATIVA", "TIPO",
                "DÍAS VACACIONES", "MES DE VACACIONES"
            ]

            for col_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=4, column=col_idx, value=h)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center_align

            current_row = 5
            for r in report_rows:
                row_vals = [
                    r['secuencial'],
                    r['cedula'],
                    r['apellidos_nombres'],
                    r['regimen'],
                    r['fecha_ingreso'],
                    r['unidad'],
                    r['tipo'],
                    r['dias_ley'],
                    r['meses']
                ]
                for col_idx, val in enumerate(row_vals, start=1):
                    cell = ws.cell(row=current_row, column=col_idx, value=val)
                    cell.font = regular_font
                    cell.border = thin_border
                    if col_idx in [1, 2, 5, 7, 8]:
                        cell.alignment = center_align
                        if col_idx == 8:
                            cell.font = bold_font
                    else:
                        cell.alignment = left_align
                ws.row_dimensions[current_row].height = 19
                current_row += 1

            # Bloque de 4 Firmas en Excel (como en la imagen de referencia)
            current_row += 2
            sig_titles_row = current_row
            ws.cell(row=sig_titles_row, column=1, value="Elaborado por:").font = bold_font
            ws.cell(row=sig_titles_row, column=4, value="Revisado por:").font = bold_font
            ws.cell(row=sig_titles_row, column=6, value="Visto bueno:").font = bold_font
            ws.cell(row=sig_titles_row, column=8, value="Aprobado por:").font = bold_font

            sig_names_row = current_row + 4  # Espacio para firma física
            ws.cell(row=sig_names_row, column=1, value=firmas['elaborado_nombre']).font = regular_font
            ws.cell(row=sig_names_row + 1, column=1, value=firmas['elaborado_cargo']).font = bold_font

            ws.cell(row=sig_names_row, column=4, value=firmas['revisado_nombre']).font = regular_font
            ws.cell(row=sig_names_row + 1, column=4, value=firmas['revisado_cargo']).font = bold_font

            ws.cell(row=sig_names_row, column=6, value=firmas['visto_bueno_nombre']).font = regular_font
            ws.cell(row=sig_names_row + 1, column=6, value=firmas['visto_bueno_cargo']).font = bold_font

            ws.cell(row=sig_names_row, column=8, value=firmas['aprobado_nombre']).font = regular_font
            ws.cell(row=sig_names_row + 1, column=8, value=firmas['aprobado_cargo']).font = bold_font

            column_widths = [6, 14, 36, 18, 15, 34, 16, 18, 28]
            for idx, width in enumerate(column_widths, start=1):
                ws.column_dimensions[get_column_letter(idx)].width = width

            buffer = BytesIO()
            wb.save(buffer)
            buffer.seek(0)
            response = HttpResponse(buffer.getvalue(),
                                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            response["Content-Disposition"] = f'attachment; filename="Consolidado_Vacaciones_{plan.year}.xlsx"'
            return response

        # 6. Exportar a PDF (Vertical)
        letterhead_path = None
        if sys_config and sys_config.letterhead:
            try:
                letterhead_path = sys_config.letterhead.path
            except Exception:
                letterhead_path = None

        template = get_template('vacation/reports/pdf_consolidated_report.html')
        html = template.render({
            'plan': plan,
            'sys_config': sys_config,
            'letterhead_path': letterhead_path,
            'report_rows': report_rows,
            'today': dt.date.today(),
            'total_funcionarios': len(report_rows),
            'firmas': firmas,
        })

        buffer = BytesIO()
        pdf = pisa.pisaDocument(BytesIO(html.encode("UTF-8")), buffer, encoding='UTF-8')
        if not pdf.err:
            response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
            response['Content-Disposition'] = f'inline; filename="Consolidado_Vacaciones_{plan.year}.pdf"'
            return response

        return HttpResponse('Error al generar el reporte consolidado en PDF', status=500)


class UserSearchSelect2View(LoginRequiredMixin, View):
    def get(self, request):
        term = (request.GET.get('q') or request.GET.get('term') or '').strip()
        if len(term) < 2:
            return JsonResponse({'results': []})

        words = term.split()
        qs = Person.objects.filter(user__isnull=False).select_related('user').distinct()

        for word in words:
            qs = qs.filter(
                Q(first_name__icontains=word) |
                Q(last_name__icontains=word) |
                Q(document_number__icontains=word) |
                Q(user__username__icontains=word)
            )

        persons = qs.order_by('last_name', 'first_name')[:25]

        results = []
        for p in persons:
            u = p.user
            nombre = (
                    getattr(u, 'signature_name', None) or
                    getattr(u, 'custom_name', None) or
                    f"{p.last_name} {p.first_name}".strip()
            )
            cargo = (
                    getattr(u, 'signature_position', None) or
                    getattr(u, 'custom_position', None) or
                    ''
            )
            text = f"{nombre.upper()}" + (f" - {cargo.upper()}" if cargo else "")
            results.append({'id': u.id, 'text': text})

        return JsonResponse({'results': results})


class VacationScheduleApprovePlanView(LoginRequiredMixin, View, PermissionRequiredMixin):
    """
    Cambia el estado a Aprobado Oficialmente (APPROVED).
    Bloquea edición, baja y el banquillo de personal.
    """

    def post(self, request, pk):
        plan = get_object_or_404(VacationSchedulePlan, pk=pk)
        if plan.status == VacationSchedulePlan.PlanStatus.APPROVED:
            return JsonResponse({
                'status': 'error',
                'success': False,
                'message': 'Este cronograma ya se encuentra aprobado oficialmente.'
            }, status=400)

        plan.status = VacationSchedulePlan.PlanStatus.APPROVED
        plan.save()

        return JsonResponse({
            'status': 'success',
            'success': True,
            'message': f'El cronograma {plan.year} ha sido Aprobado Oficialmente con éxito.'
        })


# Agregar en apps/vacation/views.py

class VacationScheduleUnitReportModalView(LoginRequiredMixin, View):
    """
    Despliega el modal para seleccionar dependencia y generar el reporte de comprobación.
    """
    template_name = 'vacation/modals/modal_unit_report.html'

    def get(self, request, plan_id):
        plan = get_object_or_404(VacationSchedulePlan, pk=plan_id)
        parent_units = AdministrativeUnit.objects.filter(
            is_active=True,
            parent__isnull=True
        ).order_by('name')

        context = {
            'plan': plan,
            'parent_units': parent_units,
        }
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            html = render_to_string(self.template_name, context, request=request)
            return HttpResponse(html)
        return render(request, self.template_name, context)


class SubUnitsAjaxView(LoginRequiredMixin, View):
    """
    Endpoint AJAX para combos dependientes de unidades administrativas.
    Si parent_id es 0, carga las unidades raíz (Nivel 1).
    Si parent_id > 0, carga las unidades hijas dependientes.
    """

    def get(self, request, parent_id=0):
        parent_id = int(parent_id or 0)

        if parent_id > 0:
            parent = get_object_or_404(AdministrativeUnit, pk=parent_id, is_active=True)
            children = AdministrativeUnit.objects.filter(parent=parent, is_active=True).order_by('name')
            units_data = [{'id': ch.id, 'name': ch.name} for ch in children]
        else:
            # Nivel 1: Direcciones principales / unidades raíz
            root_units = AdministrativeUnit.objects.filter(
                is_active=True,
                parent__isnull=True
            ).order_by('name')
            units_data = [{'id': u.id, 'name': u.name} for u in root_units]

        return JsonResponse({'units': units_data})


class VacationScheduleUnitReportPDFView(LoginRequiredMixin, View):
    """
    Genera el PDF de comprobación interna:
    Página 1: Planificados.
    Página 2: Pendientes por asignar.
    Sin firmas.
    """

    def get(self, request, plan_id):
        plan = get_object_or_404(VacationSchedulePlan, pk=plan_id)
        unit_id = request.GET.get('unit_id')
        if not unit_id:
            return HttpResponse('Debe seleccionar una unidad administrativa.', status=400)

        target_unit = get_object_or_404(AdministrativeUnit, pk=unit_id)

        # 1. Obtener jerarquía de la unidad
        unit_ids = [target_unit.id]
        descendants = [target_unit]
        while descendants:
            curr = descendants.pop(0)
            children = list(AdministrativeUnit.objects.filter(parent=curr, is_active=True))
            for ch in children:
                unit_ids.append(ch.id)
                descendants.append(ch)

        # 2. Empleados excluyendo Servicios Profesionales
        employees = Employee.objects.filter(
            is_active=True,
            area_id__in=unit_ids
        ).exclude(
            Q(employment_status__code__icontains='PROFESIONAL') |
            Q(employment_status__name__icontains='PROFESIONAL') |
            Q(employment_status__code__icontains='SERVICIOS') |
            Q(employment_status__name__icontains='SERVICIOS')
        ).select_related('person', 'employment_status').order_by('person__last_name', 'person__first_name')

        scheduled_items = VacationScheduleItem.objects.filter(
            plan=plan,
            employee__in=employees
        ).select_related('employee__person')

        # 3. Separar en asignados y pendientes
        scheduled_rows = []
        pending_rows = []

        for emp in employees:
            items = [it for it in scheduled_items if it.employee_id == emp.id]
            total_planificado = sum(it.days_quantity for it in items)

            # Límite anual
            allowed_days, ceiling, _, _ = calculate_available_days_for_month(emp, plan.year, 12)

            status_text = ((
                                   emp.employment_status.name or emp.employment_status.code or '') if emp.employment_status else '').upper()
            tipo_display = "TRABAJADOR" if any(
                k in status_text for k in ['TRABAJADOR', 'OBRERO', 'CODIGO_TRABAJO', 'CÓDIGO', 'CT']) else "EMPLEADO"
            apellidos_nombres = f"{emp.person.last_name} {emp.person.first_name}".strip()

            if total_planificado > 0:
                meses_str = ", ".join(
                    [f"{MONTH_NAMES.get(it.start_date.month, '')} ({it.days_quantity}d)" for it in items])
                scheduled_rows.append({
                    'cedula': emp.person.document_number,
                    'apellidos_nombres': apellidos_nombres,
                    'tipo': tipo_display,
                    'meses': meses_str,
                    'total_dias': total_planificado,
                })

            if total_planificado < ceiling:
                pending_rows.append({
                    'cedula': emp.person.document_number,
                    'apellidos_nombres': apellidos_nombres,
                    'tipo': tipo_display,
                    'dias_planificados': total_planificado,
                    'dias_tope': ceiling,
                    'dias_pendientes': ceiling - total_planificado,
                })

        sys_config = SystemConfiguration.get_current()
        letterhead_path = sys_config.letterhead.path if sys_config and sys_config.letterhead else None

        template = get_template('vacation/reports/pdf_unit_verification_report.html')
        html = template.render({
            'plan': plan,
            'target_unit': target_unit,
            'sys_config': sys_config,
            'letterhead_path': letterhead_path,
            'scheduled_rows': scheduled_rows,
            'pending_rows': pending_rows,
            'today': dt.date.today(),
        })

        buffer = BytesIO()
        pdf = pisa.pisaDocument(BytesIO(html.encode("UTF-8")), buffer, encoding='UTF-8')
        if not pdf.err:
            response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
            response[
                'Content-Disposition'] = f'inline; filename="Comprobacion_Vacaciones_{target_unit.id}_{plan.year}.pdf"'
            return response
        return HttpResponse('Error al generar PDF', status=500)


class VacationScheduleClearMonthView(LoginRequiredMixin, PermissionRequiredMixin, View):
    """
    Elimina todas las asignaciones planificadas en un mes específico para la dependencia activa.
    """
    permission_required = 'vacation.view_vacationscheduleplan'

    def post(self, request, plan_id, month):
        plan = get_object_or_404(VacationSchedulePlan, pk=plan_id)

        # Si el cronograma ya fue aprobado oficialmente, está bloqueado
        if plan.status == VacationSchedulePlan.PlanStatus.APPROVED:
            return JsonResponse({
                'status': 'error',
                'success': False,
                'message': 'No se pueden eliminar planificaciones en un cronograma Aprobado Oficialmente.'
            }, status=400)

        unit_id = request.GET.get('unit_id')
        user = request.user
        base_unit = None

        if unit_id and (user.is_staff or user.has_perm('vacation.change_vacationscheduleplan')):
            base_unit = AdministrativeUnit.objects.filter(pk=unit_id, is_active=True).first()
        else:
            user_emp = Employee.objects.filter(person__user=user, is_active=True).first()
            if user_emp and user_emp.area:
                base_unit = user_emp.area

        if not base_unit:
            base_unit = AdministrativeUnit.objects.filter(is_active=True, parent__isnull=True).first()

        unit_ids = []
        if base_unit:
            unit_ids = [base_unit.id]
            descendants = [base_unit]
            while descendants:
                curr = descendants.pop(0)
                children = list(AdministrativeUnit.objects.filter(parent=curr, is_active=True))
                for child in children:
                    unit_ids.append(child.id)
                    descendants.append(child)

        month_int = int(month)

        # Filtrar y eliminar solo los items de este mes y de esta dependencia
        items_to_delete = VacationScheduleItem.objects.filter(
            plan=plan,
            start_date__month=month_int,
            employee__area_id__in=unit_ids
        )
        total_deleted = items_to_delete.count()
        items_to_delete.delete()

        month_names = {
            1: 'Enero', 2: 'Febrero', 3: 'Marzo', 4: 'Abril',
            5: 'Mayo', 6: 'Junio', 7: 'Julio', 8: 'Agosto',
            9: 'Septiembre', 10: 'Octubre', 11: 'Noviembre', 12: 'Diciembre'
        }

        return JsonResponse({
            'status': 'success',
            'success': True,
            'message': f'Se vació la planificación de {month_names.get(month_int, "")}. Se liberaron {total_deleted} registros.'
        })


class ReportWrongUnitEmployeeView(LoginRequiredMixin, View):
    """Registra la novedad de que un funcionario no pertenece a la unidad actual."""

    def post(self, request, plan_id, employee_id):
        plan = get_object_or_404(VacationSchedulePlan, pk=plan_id)
        employee = get_object_or_404(Employee, pk=employee_id)
        unit_id = request.POST.get('unit_id')
        observation = request.POST.get('observation', '').strip()

        unit = AdministrativeUnit.objects.filter(pk=unit_id).first() or employee.area

        if not observation:
            return JsonResponse({'success': False, 'message': 'Debe ingresar un motivo para Talento Humano.'},
                                status=400)

        ScheduleDiscrepancyRequest.objects.create(
            plan=plan,
            unit=unit,
            employee=employee,
            request_type=ScheduleDiscrepancyRequest.RequestType.WRONG_UNIT,
            observation=observation,
            created_by=request.user
        )

        return JsonResponse({
            'success': True,
            'message': 'Se notificó la novedad a Talento Humano exitosamente.'
        })


class SearchEmployeeForScheduleView(LoginRequiredMixin, View):
    """Endpoint para el Select2 remoto con búsqueda de personal activo."""

    def get(self, request):
        term = request.GET.get('q', '').strip()
        unit_id = request.GET.get('unit_id')

        if len(term) < 2:
            return JsonResponse({'results': []})

        # Excluir personal que ya pertenezca a la unidad actual
        exclude_unit_ids = []
        if unit_id:
            _, exclude_unit_ids = resolve_scheduling_hierarchy(
                user=request.user,
                requested_unit_id=unit_id
            )

        employees = Employee.objects.select_related('person', 'area').filter(
            Q(person__document_number__icontains=term) |
            Q(person__first_name__icontains=term) |
            Q(person__last_name__icontains=term)
        ).filter(is_active=True)

        if exclude_unit_ids:
            employees = employees.exclude(area_id__in=exclude_unit_ids)

        employees = employees[:25]

        results = []
        for emp in employees:
            area_name = emp.area.name if emp.area else 'Sin Dependencia'
            full_name = f"{emp.person.last_name} {emp.person.first_name}".strip()
            results.append({
                'id': emp.id,
                'text': f"{full_name} (CI: {emp.person.document_number}) - {area_name}"
            })

        return JsonResponse({'results': results})


class RequestAddMissingEmployeeView(LoginRequiredMixin, View):
    """Registra la solicitud de incorporación de personal faltante en la unidad."""

    def post(self, request, plan_id):
        plan = get_object_or_404(VacationSchedulePlan, pk=plan_id)
        employee_id = request.POST.get('employee_id')
        unit_id = request.POST.get('unit_id')
        observation = request.POST.get('observation', '').strip()

        if not employee_id:
            return JsonResponse({'success': False, 'message': 'Debe seleccionar un servidor/trabajador.'}, status=400)

        if not observation:
            return JsonResponse({'success': False, 'message': 'Debe ingresar una justificación o mensaje.'}, status=400)

        employee = get_object_or_404(
            Employee.objects.select_related('person', 'area'),
            pk=employee_id
        )
        unit = get_object_or_404(AdministrativeUnit, pk=unit_id)

        # 1. VALIDACIÓN: ¿El servidor ya pertenece a esta dependencia o sus subdependencias?
        _, unit_ids = resolve_scheduling_hierarchy(
            user=request.user,
            requested_unit_id=unit.id
        )
        if employee.area_id in unit_ids:
            nombre = f"{employee.person.first_name} {employee.person.last_name}"
            return JsonResponse({
                'success': False,
                'message': f'El servidor {nombre} ya pertenece a esta dependencia ({employee.area.name}).'
            }, status=400)

        # 2. VALIDACIÓN: ¿Ya existe una solicitud pendiente para este servidor en esta unidad?
        already_requested = ScheduleDiscrepancyRequest.objects.filter(
            plan=plan,
            unit=unit,
            employee=employee,
            request_type=ScheduleDiscrepancyRequest.RequestType.MISSING_EMPLOYEE,
            status=ScheduleDiscrepancyRequest.Status.PENDING
        ).exists()

        if already_requested:
            return JsonResponse({
                'success': False,
                'message': f'Ya existe una solicitud pendiente enviada a Talento Humano para {employee.person.full_name}.'
            }, status=400)

        # 3. Crear solicitud
        ScheduleDiscrepancyRequest.objects.create(
            plan=plan,
            unit=unit,
            employee=employee,
            request_type=ScheduleDiscrepancyRequest.RequestType.MISSING_EMPLOYEE,
            observation=observation,
            created_by=request.user
        )

        return JsonResponse({
            'success': True,
            'message': 'Se ha enviado la solicitud de personal a Talento Humano.'
        })


class VacationScheduleItemMoveMonthView(LoginRequiredMixin, PermissionRequiredMixin, View):
    """Mueve una asignación de un mes a otro mes destino."""
    permission_required = 'vacation.view_vacationscheduleplan'

    def post(self, request, item_id):
        item = get_object_or_404(
            VacationScheduleItem.objects.select_related('plan', 'employee'),
            pk=item_id
        )
        plan = item.plan

        if plan.status == VacationSchedulePlan.PlanStatus.APPROVED:
            return JsonResponse({'success': False, 'message': 'El cronograma está aprobado y bloqueado.'}, status=400)

        new_month = request.POST.get('target_month')
        try:
            new_month = int(new_month)
            if not (1 <= new_month <= 12):
                raise ValueError()
        except (ValueError, TypeError):
            return JsonResponse({'success': False, 'message': 'Mes destino no válido.'}, status=400)

        # Si se soltó en el mismo mes donde ya estaba
        if item.start_date.month == new_month:
            return JsonResponse({'success': True, 'message': 'Sin cambios.'})

        # Verificar si el empleado ya tiene otra tarjeta en el mes destino
        already_in_target = VacationScheduleItem.objects.filter(
            plan=plan,
            employee=item.employee,
            start_date__month=new_month
        ).exclude(pk=item.id).exists()

        if already_in_target:
            return JsonResponse({
                'success': False,
                'message': f'{item.employee.person.full_name} ya tiene una planificación asignada en ese mes.'
            }, status=400)

        # Recalcular cupo permitido en el nuevo mes por si es proporcional
        allowed_int, _, _, _ = calculate_available_days_for_month(
            item.employee, plan.year, new_month
        )

        # Sumar los otros meses sin contar este item
        other_days = VacationScheduleItem.objects.filter(
            plan=plan,
            employee=item.employee
        ).exclude(pk=item.id).aggregate(total=models.Sum('days_quantity'))['total'] or 0

        disponible_en_nuevo_mes = max(0, allowed_int - other_days)

        if item.days_quantity > disponible_en_nuevo_mes:
            return JsonResponse({
                'success': False,
                'message': f'Por proporcionalidad acumulada, en ese mes solo dispone de hasta {disponible_en_nuevo_mes} días (el bloque actual tiene {item.days_quantity}d).'
            }, status=400)

        # Actualizar fechas al nuevo mes
        max_day = calendar.monthrange(plan.year, new_month)[1]
        item.start_date = dt_date(plan.year, new_month, 1)
        item.end_date = dt_date(plan.year, new_month, min(item.days_quantity, max_day))
        item.save()

        return JsonResponse({
            'success': True,
            'message': f'Se movió la asignación exitosamente.'
        })


class ScheduleDiscrepancyManageListView(LoginRequiredMixin, PermissionRequiredMixin, ListView):
    """Bandeja de gestión de novedades y solicitudes de cronograma para Talento Humano."""
    model = ScheduleDiscrepancyRequest
    template_name = 'vacation/schedule_discrepancies_list.html'
    context_object_name = 'requests_list'
    permission_required = 'vacation.change_vacationscheduleplan'

    def get_queryset(self):
        plan_id = self.kwargs.get('plan_id')
        qs = ScheduleDiscrepancyRequest.objects.filter(
            plan_id=plan_id
        ).select_related(
            'employee__person',
            'employee__area',
            'unit',
            'created_by'
        ).order_by('-created_at')

        status_filter = self.request.GET.get('status')
        if status_filter:
            qs = qs.filter(status=status_filter)
        return qs

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        plan = get_object_or_404(VacationSchedulePlan, pk=self.kwargs['plan_id'])
        context['plan'] = plan
        context['current_status'] = self.request.GET.get('status', '')

        # Contadores rápidos
        context['count_pending'] = ScheduleDiscrepancyRequest.objects.filter(plan=plan,
                                                                             status=ScheduleDiscrepancyRequest.Status.PENDING).count()
        context['count_resolved'] = ScheduleDiscrepancyRequest.objects.filter(plan=plan,
                                                                              status=ScheduleDiscrepancyRequest.Status.RESOLVED).count()
        context['count_rejected'] = ScheduleDiscrepancyRequest.objects.filter(plan=plan,
                                                                              status=ScheduleDiscrepancyRequest.Status.REJECTED).count()
        return context


class ResolveDiscrepancyRequestView(LoginRequiredMixin, PermissionRequiredMixin, View):
    permission_required = 'vacation.change_vacationscheduleplan'

    def post(self, request, pk):
        discrepancy = get_object_or_404(
            ScheduleDiscrepancyRequest.objects.select_related('employee', 'unit', 'plan'),
            pk=pk
        )
        action = request.POST.get('action')  # 'resolve' o 'reject'
        admin_note = request.POST.get('admin_note', '').strip()
        target_unit_id = request.POST.get('target_unit_id')

        if action == 'resolve':
            emp = discrepancy.employee
            # Determinar a qué unidad reubicar
            target_unit = None
            if target_unit_id:
                target_unit = AdministrativeUnit.objects.filter(pk=target_unit_id).first()
            elif discrepancy.request_type == ScheduleDiscrepancyRequest.RequestType.MISSING_EMPLOYEE:
                target_unit = discrepancy.unit

            if target_unit:
                emp.area = target_unit
                emp.save(update_fields=['area'])

            discrepancy.status = ScheduleDiscrepancyRequest.Status.RESOLVED
            nota_th = f"\n[TH Resuelto]: Reubicado a {emp.area.name}."
            if admin_note:
                nota_th += f" Nota: {admin_note}"
            discrepancy.observation += nota_th
            discrepancy.save()

            return JsonResponse({
                'success': True,
                'message': f'Novedad resuelta con éxito. {emp.person.full_name} reubicado a {emp.area.name}.'
            })

        elif action == 'reject':
            discrepancy.status = ScheduleDiscrepancyRequest.Status.REJECTED
            if admin_note:
                discrepancy.observation += f"\n[TH Rechazado]: {admin_note}"
            discrepancy.save()

            return JsonResponse({
                'success': True,
                'message': 'La novedad ha sido desestimada.'
            })

        return JsonResponse({'success': False, 'message': 'Acción no válida.'}, status=400)
