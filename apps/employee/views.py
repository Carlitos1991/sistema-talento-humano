import json
import logging
from datetime import date, time, datetime
from decimal import Decimal
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from django.contrib.auth import get_user_model
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin, PermissionRequiredMixin
from django.core.exceptions import ObjectDoesNotExist
from django.db import transaction
from django.http import JsonResponse, Http404, HttpResponse
from django.shortcuts import get_object_or_404, render
from django.template.loader import get_template
from django.urls import reverse
from django.utils import timezone
from django.utils.decorators import method_decorator
from django.views import View
from django.views.decorators.http import require_POST
from django.views.generic import DetailView, ListView, CreateView, UpdateView
from xhtml2pdf import pisa

from biometric.models import OfflineAttendanceRegistry
from budget.models import BudgetAssignmentHistory, BudgetLine
from contract.models import ManagementPeriod
from core.models import CatalogItem, Location, SystemConfiguration
from institution.models import AdministrativeUnit
from payroll.models import Payslip
from permitrequest.models import PermitRequest, PermitType
from person.models import Person, PersonAuditLog
from person.utils import log_person_audit, PERSON_AUDIT_SECTIONS
from personnel_actions.models import PersonnelAction
from sanctions.models import Sanction
from schedule.models import EmployeeScheduleHistory
from vacation.models import EmployeeVacationBalance

from .forms import (
    AcademicTitleForm, WorkExperienceForm, TrainingForm,
    BankAccountForm, PayrollInfoForm, InstitutionalDataForm
)
from .models import (
    Employee, Curriculum, AcademicTitle, WorkExperience, Training,
    InstitutionalData, EconomicData, EmployeeProfileVisibility, TeleworkActivity
)

User = get_user_model()
logger = logging.getLogger(__name__)


def _safe_related(instance, attr_name, default=None):
    if instance is None:
        return default
    try:
        return getattr(instance, attr_name)
    except (ObjectDoesNotExist, Exception):
        return default


class PersonAuditModalView(LoginRequiredMixin, PermissionRequiredMixin, View):
    permission_required = 'person.view_person'
    template_name = 'person/modals/modal_person_audit.html'

    def get(self, request, person_id):
        person = get_object_or_404(Person, pk=person_id)
        today = timezone.now().date()
        today_str = today.isoformat()

        start_date_str = request.GET.get('start_date')
        end_date_str = request.GET.get('end_date')
        show_all = request.GET.get('show_all') == '1'

        # Ordenamiento estricto: del más reciente al más antiguo (-created_at, -id)
        logs_qs = PersonAuditLog.objects.filter(person=person).select_related('user').order_by('-created_at', '-id')

        if show_all:
            start_date_str = ''
            end_date_str = ''
            audit_logs = logs_qs[:500]
        else:
            if not start_date_str:
                start_date_str = today_str
            if not end_date_str:
                end_date_str = today_str

            try:
                start_d = datetime.strptime(start_date_str, '%Y-%m-%d').date()
                end_d = datetime.strptime(end_date_str, '%Y-%m-%d').date()
                audit_logs = logs_qs.filter(created_at__date__range=(start_d, end_d))[:1000]
            except ValueError:
                audit_logs = logs_qs.filter(created_at__date=today)[:500]

        return render(request, self.template_name, {
            'person': person,
            'audit_logs': audit_logs,
            'start_date': start_date_str,
            'end_date': end_date_str,
            'is_all': show_all,
            'today': today_str,
        })


@login_required
def search_employee_by_cedula(request):
    cedula = request.GET.get('q', '').strip()
    if not cedula:
        return JsonResponse({'success': False, 'message': 'Cédula no proporcionada.'})

    try:
        emp = Employee.objects.select_related('person').get(person__document_number=cedula)
        existing_assignment = BudgetLine.objects.filter(current_employee=emp).first()
        if existing_assignment:
            return JsonResponse({
                'success': False,
                'message': f'La persona {emp.person.full_name} ya tiene asignada la partida {existing_assignment.code}.'
            })
        return JsonResponse({
            'success': True,
            'id': emp.id,
            'full_name': emp.person.full_name,
            'email': emp.person.email or 'Sin correo registrado',
            'photo_url': emp.person.photo.url if emp.person.photo else None
        })
    except Employee.DoesNotExist:
        try:
            p = Person.objects.get(document_number=cedula)
            if hasattr(p, 'employee_profile') and p.employee_profile:
                emp = p.employee_profile
                existing_assignment = BudgetLine.objects.filter(current_employee=emp).first()
                if existing_assignment:
                    return JsonResponse({
                        'success': False,
                        'message': f'La persona {emp.person.full_name} ya tiene asignada la partida {existing_assignment.code}.'
                    })
                return JsonResponse({
                    'success': True,
                    'id': emp.id,
                    'full_name': emp.person.full_name,
                    'email': emp.person.email or 'Sin correo registrado',
                    'photo_url': emp.person.photo.url if emp.person.photo else None
                })
            return JsonResponse({
                'success': False,
                'message': 'Se encontró la Persona pero no tiene perfil de Empleado.'
            })
        except Person.DoesNotExist:
            return JsonResponse({'success': False, 'message': 'No se encontró registro con esa cédula.'})


@require_POST
@login_required
def log_tab_audit_api(request, person_id):
    person = get_object_or_404(Person, pk=person_id)
    try:
        data = json.loads(request.body)
    except Exception:
        data = request.POST

    tab_name = data.get('tab_name', '').strip()

    # Si no viene un nombre de pestaña válido, no registrar nada
    if not tab_name or tab_name.lower() == 'sección':
        return JsonResponse({'success': False, 'message': 'Pestaña no especificada.'}, status=400)

    detail_text = f'Revisó {tab_name}'

    # PREVENCIÓN DE DUPLICADOS: Verificar si se acaba de registrar exactamente lo mismo hace menos de 2 segundos
    last_log = PersonAuditLog.objects.filter(
        person=person,
        user=request.user
    ).order_by('-created_at', '-id').first()

    if last_log and last_log.action_detail == detail_text:
        time_diff = (timezone.now() - last_log.created_at).total_seconds()
        if time_diff < 2:
            return JsonResponse({'success': True, 'message': 'Registro omitido por duplicidad.'})

    tab_id = data.get('tab_id', 'general')
    section = PERSON_AUDIT_SECTIONS.get(tab_id, tab_id) if isinstance(PERSON_AUDIT_SECTIONS, dict) else tab_id

    log_person_audit(
        request,
        person,
        PersonAuditLog.Action.VIEW,
        section,
        detail_text
    )
    return JsonResponse({'success': True})


@require_POST
@login_required
def upload_person_photo_api(request, person_id):
    person = get_object_or_404(Person, pk=person_id)
    photo_file = request.FILES.get('photo')
    if not photo_file:
        return JsonResponse({'success': False, 'message': 'No se ha proporcionado ninguna imagen.'}, status=400)

    person.photo = photo_file
    person.save(update_fields=['photo'])

    # REGISTRO EN AUDITORÍA
    log_person_audit(
        request,
        person,
        PersonAuditLog.Action.UPDATE,
        PERSON_AUDIT_SECTIONS.get('personal', 'personal') if isinstance(PERSON_AUDIT_SECTIONS, dict) else 'personal',
        'Actualizó foto de perfil'
    )

    return JsonResponse({
        'success': True,
        'message': 'Foto de perfil actualizada correctamente.',
        'photo_url': person.photo.url
    })


class EmployeeDetailWizardView(LoginRequiredMixin, PermissionRequiredMixin, DetailView):
    model = Person
    template_name = 'employee/employee_detail_wizard.html'
    context_object_name = 'person'
    permission_required = 'person.view_person'

    def has_permission(self):
        if self.request.user.is_staff or self.request.user.is_superuser:
            return True
        return super().has_permission()

    def get_queryset(self):
        return Person.objects.select_related(
            'employee_profile__area',
            'employee_profile__employment_status',
            'curriculum',
            'economic_data__bank_account',
            'economic_data__payroll_info',
            'document_type', 'gender', 'country', 'province', 'canton', 'parish'
        ).all()

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        user_person = _safe_related(self.request.user, 'person', None)
        employee = _safe_related(self.object, 'employee_profile', None)

        curriculum, _ = Curriculum.objects.get_or_create(person=self.object)
        economic_data, _ = EconomicData.objects.get_or_create(person=self.object)
        institutional_data = None
        current_schedule_assignment = None

        if employee:
            institutional_data, _ = InstitutionalData.objects.get_or_create(employee=employee)
            current_schedule_assignment = EmployeeScheduleHistory.objects.filter(
                employee=employee, is_current=True
            ).select_related('schedule').first()

        context['current_schedule_assignment'] = current_schedule_assignment
        context['curriculum_titles_count'] = 0
        context['curriculum_experiences_count'] = 0
        context['curriculum_courses_count'] = 0
        context['employee_area_name'] = employee.area.name if (employee and employee.area) else 'SIN AREA ASIGNADA'
        context['employee_profile'] = employee
        context['curriculum_obj'] = curriculum
        context['areas_list'] = AdministrativeUnit.objects.filter(is_active=True).order_by('name')
        context['economic_data'] = economic_data
        context['bank_account'] = _safe_related(economic_data, 'bank_account', None)
        context['payroll_info'] = _safe_related(economic_data, 'payroll_info', None)
        context['institutional_data'] = institutional_data

        if curriculum:
            try:
                context['curriculum_titles_count'] = curriculum.academic_titles.count()
                context['curriculum_experiences_count'] = curriculum.work_experiences.count()
                context['curriculum_courses_count'] = curriculum.trainings.count()
                total_days = 0
                for exp in curriculum.work_experiences.all():
                    start = exp.start_date
                    end = date.today() if exp.is_current else (exp.end_date or date.today())
                    if start:
                        total_days += (end - start).days

                context['experience_years'] = total_days // 365
                context['experience_months'] = (total_days % 365) // 30
            except Exception:
                context['experience_years'] = 0
                context['experience_months'] = 0

            log_person_audit(
                self.request,
                self.object,
                PersonAuditLog.Action.VIEW,
                'general',
                'Visualizó Expediente Digital'
            )

        # Visibilidad de pestañas
        default_enabled_tabs = ['personal', 'curriculum', 'economic', 'institutional']
        visibilities = []
        if hasattr(self.object, 'user') and self.object.user:
            visibilities = EmployeeProfileVisibility.objects.filter(
                user=self.object.user
            ).values_list('tab_id', 'is_visible')

        visibility_dict = {v[0]: v[1] for v in visibilities}
        all_tabs = [
            'personal', 'curriculum', 'economic', 'institutional',
            'budget', 'contracts', 'actions', 'permissions',
            'payments', 'sanctions', 'vacations', 'schedule', 'telework'
        ]

        final_visibility = {}
        for tab_id in all_tabs:
            if tab_id in visibility_dict:
                final_visibility[tab_id] = str(visibility_dict[tab_id]).lower()
            else:
                final_visibility[tab_id] = str(tab_id in default_enabled_tabs).lower()

        context['tab_visibilities'] = json.dumps(final_visibility)
        context['can_edit_person'] = self.request.user.has_perm('person.change_person')

        # Jerarquía institucional
        hierarchy_list = []
        if employee and employee.area:
            unit = employee.area
            while unit:
                hierarchy_list.insert(0, {'name': unit.name, 'level_name': unit.level.name})
                unit = unit.parent
        context['hierarchy_list'] = hierarchy_list

        # Partida presupuestaria
        try:
            if employee:
                current = BudgetLine.objects.filter(current_employee=employee).first()
                context['current_partida'] = {
                    'code': current.number_individual,
                    'budget': current.code,
                    'name': (current.position_item.name if current.position_item else '') or str(current),
                    'remuneration': str(current.remuneration) if current.remuneration is not None else '—',
                    'category': (current.category_item.name if getattr(current, 'category_item', None) else '')
                } if current else None

                assignments = BudgetAssignmentHistory.objects.filter(employee=employee).select_related('budget_line')
                context['partida_history'] = [{
                    'partida_code': a.budget_line.number_individual or a.budget_line.code,
                    'position_name': (a.budget_line.position_item.name if a.budget_line.position_item else ''),
                    'remuneration': str(a.budget_line.remuneration),
                    'partida_name': str(a.budget_line),
                    'code': a.budget_line.code
                } for a in assignments]
            else:
                context['current_partida'] = None
                context['partida_history'] = []
        except Exception:
            context['current_partida'] = None
            context['partida_history'] = []

        # Contratos
        try:
            if employee:
                latest = ManagementPeriod.objects.filter(employee=employee).order_by('-start_date').first()
                context['current_contract'] = {
                    'id': latest.id,
                    'document_number': latest.document_number,
                    'position_name': latest.get_dynamic_position,
                    'contract_type_name': str(latest.contract_type) if latest.contract_type else '',
                    'start_date': latest.start_date.strftime('%d/%m/%Y') if latest.start_date else '',
                    'end_date': latest.end_date.strftime('%d/%m/%Y') if latest.end_date else '',
                    'administrative_unit': latest.administrative_unit.name if latest.administrative_unit else '',
                    'status_name': latest.status.name if latest.status else '',
                    'signed_document_url': latest.signed_document.url if getattr(latest, 'signed_document',
                                                                                 None) else None
                } if latest else None

                periods = ManagementPeriod.objects.filter(employee=employee).select_related(
                    'contract_type', 'status', 'administrative_unit'
                ).order_by('-start_date')[:500]
                context['contract_history'] = [{
                    'id': per.id,
                    'document_number': per.document_number,
                    'position_name': per.get_dynamic_position,
                    'contract_type_name': str(per.contract_type) if per.contract_type else '',
                    'start_date': per.start_date.strftime('%d/%m/%Y') if per.start_date else '',
                    'end_date': per.end_date.strftime('%d/%m/%Y') if per.end_date else '',
                    'administrative_unit': per.administrative_unit.name if per.administrative_unit else '',
                    'status_name': per.status.name if per.status else '',
                    'signed_document_url': per.signed_document.url if getattr(per, 'signed_document', None) else None
                } for per in periods]
            else:
                context['current_contract'] = None
                context['contract_history'] = []
        except Exception:
            context['current_contract'] = None
            context['contract_history'] = []

        # Historial de acciones de personal
        try:
            if employee:
                actions_qs = PersonnelAction.objects.filter(employee=employee).select_related('action_type').order_by(
                    '-date_issue')[:500]
                context['actions_list'] = [{
                    'id': a.pk,
                    'number': a.number,
                    'action_name': a.action_type.name if a.action_type else '',
                    'from_area': getattr(getattr(a, 'movement', None), 'previous_unit', '') or '',
                    'to_area': getattr(getattr(a, 'movement', None), 'new_unit', '') or '',
                    'issued_date': a.date_issue,
                    'effective_date': a.date_effective,
                    'document_url': reverse('personnel_actions:action_pdf', args=[a.pk])
                } for a in actions_qs]
            else:
                context['actions_list'] = []
        except Exception:
            context['actions_list'] = []

        # Roles de pago
        try:
            if employee:
                payslips_qs = Payslip.objects.filter(employee=employee).select_related('period').order_by(
                    '-period__year', '-period__id')[:12]
                context['payment_roles_history'] = [{
                    'id': payslip.pk,
                    'period_month': payslip.period.month if payslip.period else '',
                    'period_year': payslip.period.year if payslip.period else '',
                    'total_income': float(payslip.total_income or 0),
                    'total_deduction': float(payslip.total_deduction or 0),
                    'net_pay': float(payslip.net_pay or 0),
                    'print_url': reverse('payroll:payslip_print', args=[payslip.pk])
                } for payslip in payslips_qs]
            else:
                context['payment_roles_history'] = []
        except Exception:
            context['payment_roles_history'] = []

        # Historial de sanciones
        try:
            if employee:
                sanctions_qs = Sanction.objects.filter(employee=employee).select_related('sanction_type').order_by(
                    '-sanction_date')[:500]
                context['sanctions_history'] = [{
                    'type': s.sanction_type.name if s.sanction_type else '',
                    'description': s.description,
                    'reason': s.legal_basis,
                    'severity': dict(Sanction.SEVERITY_CHOICES).get(s.severity, s.severity) if hasattr(Sanction,
                                                                                                       'SEVERITY_CHOICES') else s.severity,
                    'severity_code': s.severity,
                    'date': s.sanction_date
                } for s in sanctions_qs]
            else:
                context['sanctions_history'] = []
        except Exception:
            context['sanctions_history'] = []

        # Balances de vacaciones
        try:
            if employee:
                balances_qs = EmployeeVacationBalance.objects.filter(employee=employee).select_related(
                    'period').order_by('-created_at')[:500]
                context['vacation_balances'] = [{
                    'id': b.id,
                    'period': b.period.name if b.period else '',
                    'total_days': float(b.total_days or 0),
                    'total_with_previous_balance': float(
                        (Decimal(str(b.total_days or 0)) + Decimal(str(b.additional_days or 0)))),
                    'balance_days': float(b.balance_days or 0),
                    'additional_days': float(b.additional_days or 0),
                    'permit_days': float(b.permit_days or 0),
                    'vacation_days': float(b.vacation_days or 0),
                    'taken_days': float(getattr(b, 'taken_days', 0) or 0),
                    'observation': b.observation
                } for b in balances_qs]

                last_balance = balances_qs.first()
                if last_balance:
                    total_cap = Decimal(str(last_balance.total_days or 0)) + Decimal(
                        str(last_balance.additional_days or 0))
                    p_used = Decimal(str(last_balance.permit_days or 0))
                    v_used = Decimal(str(last_balance.vacation_days or 0))
                    s_db = Decimal(str(last_balance.balance_days or (total_cap - (p_used + v_used))))
                else:
                    total_cap, p_used, v_used, s_db = Decimal('0.0'), Decimal('0.0'), Decimal('0.0'), Decimal('0.0')

                context['vacation_chart'] = {
                    'total_capacity': float(total_cap),
                    'permits': float(p_used),
                    'vacations': float(v_used),
                    'saldo': float(s_db)
                }
            else:
                context['vacation_balances'] = []
                context['vacation_chart'] = {'total_capacity': 0, 'permits': 0, 'vacations': 0, 'saldo': 0}
        except Exception:
            context['vacation_balances'] = []
            context['vacation_chart'] = {'total_capacity': 0, 'permits': 0, 'vacations': 0, 'saldo': 0}

        return context


class EmployeeSelfDashboardView(EmployeeDetailWizardView):
    template_name = 'employee/employee_dashboard.html'
    permission_required = ()

    def get_object(self, queryset=None):
        user_person = _safe_related(self.request.user, 'person', None)
        if not user_person:
            raise Http404("El usuario no tiene persona asociada.")
        self.is_self_dashboard = True
        return user_person


@transaction.atomic
def upload_cv_pdf(request, person_id):
    if request.method == 'POST':
        person = get_object_or_404(Person, pk=person_id)
        curriculum, _ = Curriculum.objects.get_or_create(person=person)
        pdf_file = request.FILES.get('pdf_file')
        if pdf_file:
            curriculum.pdf_file = pdf_file
            curriculum.save()
            return JsonResponse({'success': True, 'message': 'PDF actualizado correctamente.'})
    return JsonResponse({'success': False, 'message': 'Error al subir archivo.'}, status=400)


# =====================================================================
# VISTAS MODALES CRUD (TITULOS, EXPERIENCIA, CAPACITACIONES)
# =====================================================================

class AcademicTitleModalListView(LoginRequiredMixin, PermissionRequiredMixin, ListView):
    model = AcademicTitle
    template_name = 'employee/modals/modal_academic_title_list.html'
    context_object_name = 'titles'
    permission_required = 'employee.view_academictitle'

    def get_queryset(self):
        self.person = get_object_or_404(Person, pk=self.kwargs['person_id'])
        return AcademicTitle.objects.filter(curriculum__person=self.person).select_related('education_level').order_by(
            '-graduation_year', '-pk')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['person'] = self.person
        return context


class AcademicTitleCreateView(LoginRequiredMixin, PermissionRequiredMixin, CreateView):
    model = AcademicTitle
    form_class = AcademicTitleForm
    template_name = 'employee/modals/modal_academic_title.html'
    permission_required = 'employee.add_academictitle'

    def dispatch(self, request, *args, **kwargs):
        self.person = get_object_or_404(Person, pk=self.kwargs['person_id'])
        return super().dispatch(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['person'] = self.person
        return context

    def form_valid(self, form):
        title = form.save(commit=False)
        curriculum, _ = Curriculum.objects.get_or_create(person=self.person)
        title.curriculum = curriculum
        title.save()
        return JsonResponse({'success': True, 'message': 'Título registrado exitosamente.'})

    def form_invalid(self, form):
        return render(self.request, self.template_name, self.get_context_data(form=form), status=400)


class AcademicTitleUpdateView(LoginRequiredMixin, PermissionRequiredMixin, UpdateView):
    model = AcademicTitle
    form_class = AcademicTitleForm
    template_name = 'employee/modals/modal_academic_title.html'
    permission_required = 'employee.change_academictitle'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['person'] = self.object.curriculum.person
        return context

    def form_valid(self, form):
        form.save()
        return JsonResponse({'success': True, 'message': 'Título actualizado correctamente.'})

    def form_invalid(self, form):
        return render(self.request, self.template_name, self.get_context_data(form=form), status=400)


class AcademicTitleDeleteView(LoginRequiredMixin, PermissionRequiredMixin, View):
    permission_required = 'employee.delete_academictitle'

    def post(self, request, pk):
        title = get_object_or_404(AcademicTitle, pk=pk)
        title.delete()
        return JsonResponse({'success': True, 'message': 'Título eliminado correctamente.'})


class WorkExperienceModalListView(LoginRequiredMixin, PermissionRequiredMixin, ListView):
    model = WorkExperience
    template_name = 'employee/modals/modal_work_experience_list.html'
    context_object_name = 'experiences'
    permission_required = 'person.view_person'

    def get_queryset(self):
        self.person = get_object_or_404(Person, pk=self.kwargs['person_id'])
        return WorkExperience.objects.filter(curriculum__person=self.person).order_by('-start_date', '-pk')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['person'] = self.person
        return context


class WorkExperienceCreateView(LoginRequiredMixin, PermissionRequiredMixin, CreateView):
    model = WorkExperience
    form_class = WorkExperienceForm
    template_name = 'employee/modals/modal_work_experience.html'
    permission_required = 'person.change_person'

    def dispatch(self, request, *args, **kwargs):
        self.person = get_object_or_404(Person, pk=self.kwargs['person_id'])
        return super().dispatch(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['person'] = self.person
        return context

    def form_valid(self, form):
        experience = form.save(commit=False)
        curriculum, _ = Curriculum.objects.get_or_create(person=self.person)
        experience.curriculum = curriculum
        experience.save()
        return JsonResponse({'success': True, 'message': 'Experiencia registrada exitosamente.'})

    def form_invalid(self, form):
        return render(self.request, self.template_name, self.get_context_data(form=form), status=400)


class WorkExperienceUpdateView(LoginRequiredMixin, PermissionRequiredMixin, UpdateView):
    model = WorkExperience
    form_class = WorkExperienceForm
    template_name = 'employee/modals/modal_work_experience.html'
    permission_required = 'person.change_person'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        curriculum = getattr(self.object, 'curriculum', None)
        context['person'] = getattr(curriculum, 'person', None)
        return context

    def form_valid(self, form):
        form.save()
        return JsonResponse({'success': True, 'message': 'Experiencia actualizada correctamente.'})

    def form_invalid(self, form):
        return render(self.request, self.template_name, self.get_context_data(form=form), status=400)


class WorkExperienceDeleteView(LoginRequiredMixin, PermissionRequiredMixin, View):
    permission_required = 'person.change_person'

    def post(self, request, pk):
        exp = get_object_or_404(WorkExperience, pk=pk)
        exp.delete()
        return JsonResponse({'success': True, 'message': 'Experiencia eliminada correctamente.'})


class CoursesModalListView(LoginRequiredMixin, PermissionRequiredMixin, ListView):
    model = Training
    template_name = 'employee/modals/modal_training_list.html'
    context_object_name = 'courses'
    permission_required = 'person.change_person'

    def get_queryset(self):
        self.person = get_object_or_404(Person, pk=self.kwargs['person_id'])
        return Training.objects.filter(curriculum__person=self.person).order_by('-completion_date', '-pk')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['person'] = self.person
        return context


class CoursesCreateView(LoginRequiredMixin, PermissionRequiredMixin, CreateView):
    model = Training
    form_class = TrainingForm
    template_name = 'employee/modals/modal_training.html'
    permission_required = 'person.change_person'

    def dispatch(self, request, *args, **kwargs):
        self.person = get_object_or_404(Person, pk=self.kwargs['person_id'])
        return super().dispatch(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['person'] = self.person
        return context

    def form_valid(self, form):
        training = form.save(commit=False)
        curriculum, _ = Curriculum.objects.get_or_create(person=self.person)
        training.curriculum = curriculum
        training.save()
        return JsonResponse({'success': True, 'message': 'Capacitación registrada exitosamente.'})

    def form_invalid(self, form):
        return render(self.request, self.template_name, self.get_context_data(form=form), status=400)


class CoursesUpdateView(LoginRequiredMixin, PermissionRequiredMixin, UpdateView):
    model = Training
    form_class = TrainingForm
    template_name = 'employee/modals/modal_training.html'
    permission_required = 'person.change_person'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        curriculum = getattr(self.object, 'curriculum', None)
        context['person'] = getattr(curriculum, 'person', None)
        return context

    def form_valid(self, form):
        form.save()
        return JsonResponse({'success': True, 'message': 'Capacitación actualizada correctamente.'})

    def form_invalid(self, form):
        return render(self.request, self.template_name, self.get_context_data(form=form), status=400)


class CoursesDeleteView(LoginRequiredMixin, PermissionRequiredMixin, View):
    permission_required = 'person.change_person'

    def post(self, request, pk):
        training = get_object_or_404(Training, pk=pk)
        training.delete()
        return JsonResponse({'success': True, 'message': 'Capacitación eliminada correctamente.'})


# =====================================================================
# MODALES DE DATOS INSTITUCIONALES Y ECONÓMICOS
# =====================================================================

class InstitutionalDataUpdateView(LoginRequiredMixin, PermissionRequiredMixin, View):
    permission_required = 'person.change_person'
    template_name = 'employee/modals/modal_institutional_data_form.html'

    def get(self, request, person_id):
        person = get_object_or_404(Person, pk=person_id)
        employee = getattr(person, 'employee_profile', None)
        if not employee:
            return JsonResponse({'success': False, 'message': 'El empleado no tiene perfil registrado.'}, status=404)

        inst_data, _ = InstitutionalData.objects.get_or_create(employee=employee)
        form = InstitutionalDataForm(instance=inst_data)
        return render(request, self.template_name, {'form': form, 'person': person, 'institutional_data': inst_data})

    def post(self, request, person_id):
        person = get_object_or_404(Person, pk=person_id)
        employee = getattr(person, 'employee_profile', None)
        if not employee:
            return JsonResponse({'success': False, 'message': 'El empleado no tiene perfil registrado.'}, status=404)

        inst_data, _ = InstitutionalData.objects.get_or_create(employee=employee)
        form = InstitutionalDataForm(request.POST, instance=inst_data)
        if form.is_valid():
            inst = form.save(commit=False)
            if not inst.institutional_email:
                inst.institutional_email = None
            inst.save()
            log_person_audit(request, person, PersonAuditLog.Action.UPDATE, PERSON_AUDIT_SECTIONS['institutional'],
                             'Actualizó datos institucionales')
            return JsonResponse({'success': True, 'message': 'Datos institucionales actualizados correctamente.'})
        return render(request, self.template_name, {'form': form, 'person': person, 'institutional_data': inst_data},
                      status=400)


class BankAccountUpdateView(LoginRequiredMixin, PermissionRequiredMixin, View):
    permission_required = 'person.change_person'
    template_name = 'employee/modals/modal_bank_account_form.html'

    def get(self, request, person_id):
        person = get_object_or_404(Person, pk=person_id)
        economic_data, _ = EconomicData.objects.get_or_create(person=person)
        instance = getattr(economic_data, 'bank_account', None)
        form = BankAccountForm(instance=instance)
        return render(request, self.template_name, {'form': form, 'person': person})

    def post(self, request, person_id):
        person = get_object_or_404(Person, pk=person_id)
        economic_data, _ = EconomicData.objects.get_or_create(person=person)
        instance = getattr(economic_data, 'bank_account', None)
        form = BankAccountForm(request.POST, instance=instance)
        if form.is_valid():
            bank_acc = form.save(commit=False)
            bank_acc.economic_data = economic_data
            bank_acc.save()
            log_person_audit(request, person, PersonAuditLog.Action.UPDATE, PERSON_AUDIT_SECTIONS['economic'],
                             'Actualizó cuenta bancaria')
            return JsonResponse({'success': True, 'message': 'Cuenta bancaria guardada correctamente.'})
        return render(request, self.template_name, {'form': form, 'person': person}, status=400)


class PayrollInfoUpdateView(LoginRequiredMixin, PermissionRequiredMixin, View):
    permission_required = 'person.change_person'
    template_name = 'employee/modals/modal_payroll_info_form.html'

    def get(self, request, person_id):
        person = get_object_or_404(Person, pk=person_id)
        economic_data, _ = EconomicData.objects.get_or_create(person=person)
        instance = getattr(economic_data, 'payroll_info', None)
        form = PayrollInfoForm(instance=instance)
        return render(request, self.template_name, {'form': form, 'person': person})

    def post(self, request, person_id):
        person = get_object_or_404(Person, pk=person_id)
        economic_data, _ = EconomicData.objects.get_or_create(person=person)
        instance = getattr(economic_data, 'payroll_info', None)
        form = PayrollInfoForm(request.POST, instance=instance)
        if form.is_valid():
            payroll = form.save(commit=False)
            payroll.economic_data = economic_data
            payroll.save()
            log_person_audit(request, person, PersonAuditLog.Action.UPDATE, PERSON_AUDIT_SECTIONS['economic'],
                             'Actualizó información de nómina')
            return JsonResponse({'success': True, 'message': 'Información de nómina actualizada correctamente.'})
        return render(request, self.template_name, {'form': form, 'person': person}, status=400)


class ContractDetailModalView(LoginRequiredMixin, PermissionRequiredMixin, DetailView):
    model = ManagementPeriod
    template_name = 'employee/modals/modal_contract_detail.html'
    context_object_name = 'contract'
    permission_required = 'person.view_person'

    def get_queryset(self):
        return ManagementPeriod.objects.select_related('contract_type', 'status', 'administrative_unit',
                                                       'employee__person')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['person'] = self.object.employee.person
        return context


# =====================================================================
# TELETRABAJO
# =====================================================================

@login_required
def telework_activity_modal_view(request, person_id):
    person = get_object_or_404(Person, pk=person_id)
    return render(request, 'employee/modals/modal_activity_telework_form.html', {'person': person})


@login_required
def telework_report_modal_view(request, person_id):
    person = get_object_or_404(Person, pk=person_id)
    today = timezone.now().date()
    first_day = today.replace(day=1)
    return render(request, 'employee/modals/modal_telework_report.html', {
        'person': person,
        'today': today.isoformat(),
        'first_day': first_day.isoformat(),
    })


@login_required
def get_telework_data_api(request, person_id):
    person = get_object_or_404(Person, pk=person_id)
    employee = person.employee_profile
    today = timezone.now().date()
    is_own_profile = (request.user == person.user)

    punches = OfflineAttendanceRegistry.objects.filter(employee=employee, captured_at__date=today).order_by(
        '-captured_at')
    last_punch = punches.first()
    last_punch_type = last_punch.punch_type if last_punch else None
    has_income_today = OfflineAttendanceRegistry.objects.filter(employee=employee, captured_at__date=today,
                                                                punch_type='INCOME').exists()
    activities = TeleworkActivity.objects.filter(employee=employee, created_at__date=today).order_by('-created_at')

    return JsonResponse({
        'success': True,
        'is_own_profile': is_own_profile,
        'last_punch_type': last_punch_type,
        'has_income': has_income_today,
        'punches': [{
            'type_code': p.punch_type,
            'type': p.get_punch_type_display(),
            'time': p.captured_at.strftime('%H:%M:%S'),
        } for p in punches],
        'activities': [{
            'id': a.id,
            'title': a.title,
            'detail': a.detail,
            'percentage': a.percentage,
            'time': a.created_at.strftime('%H:%M')
        } for a in activities]
    })


@require_POST
def add_telework_activity_api(request, person_id):
    person = get_object_or_404(Person, pk=person_id)
    if request.user != person.user:
        return JsonResponse({'success': False, 'message': 'No autorizado para registrar actividades en este perfil.'},
                            status=403)

    employee = person.employee_profile
    today = timezone.now().date()
    last_punch = OfflineAttendanceRegistry.objects.filter(employee=employee, captured_at__date=today).order_by(
        '-captured_at').first()

    if not last_punch or last_punch.punch_type == 'EXIT':
        return JsonResponse({'success': False, 'message': 'Debe tener una marcación de ENTRADA activa para este día.'},
                            status=400)

    title = request.POST.get('title')
    if not title:
        return JsonResponse({'success': False, 'message': 'El título de la actividad es obligatorio.'}, status=400)

    try:
        TeleworkActivity.objects.create(
            employee=employee,
            title=title,
            detail=request.POST.get('detail', ''),
            percentage=request.POST.get('percentage', 0),
            status='IN_PROGRESS'
        )
        return JsonResponse({'success': True, 'message': 'Actividad registrada'})
    except Exception as e:
        logger.error(f"Error saving telework activity: {e}")
        return JsonResponse({'success': False, 'message': 'Error interno al guardar la actividad.'}, status=500)


@require_POST
def mark_telework_attendance_api(request, person_id):
    person = get_object_or_404(Person, pk=person_id)
    if request.user != person.user:
        return JsonResponse({'success': False, 'message': 'No autorizado. Solo el titular puede marcar asistencia.'},
                            status=403)

    employee = person.employee_profile
    punch_type = request.POST.get('punch_type')
    today = timezone.now().date()
    last_punch = OfflineAttendanceRegistry.objects.filter(employee=employee, captured_at__date=today).order_by(
        '-captured_at').first()

    if last_punch and last_punch.punch_type == punch_type:
        tipo_str = "ENTRADA" if punch_type == 'INCOME' else "SALIDA"
        return JsonResponse({'success': False, 'message': f'Ya existe una {tipo_str} registrada como último evento.'},
                            status=400)

    if punch_type == 'EXIT':
        has_income = OfflineAttendanceRegistry.objects.filter(employee=employee, punch_type='INCOME',
                                                              captured_at__date=today).exists()
        if not has_income:
            return JsonResponse(
                {'success': False, 'message': 'No puede marcar SALIDA sin haber registrado un INGRESO previo.'},
                status=400)

    OfflineAttendanceRegistry.objects.create(
        employee=employee, punch_type=punch_type, captured_at=timezone.now(),
        latitude=request.POST.get('latitude', 0), longitude=request.POST.get('longitude', 0),
        source='WEB', sync_status='SYNCED'
    )
    return JsonResponse({'success': True, 'message': 'Marcación registrada correctamente.'})


@login_required
def generate_telework_report_pdf(request, person_id):
    person = get_object_or_404(Person, pk=person_id)
    employee = getattr(person, 'employee_profile', None)

    start_date_str = request.GET.get('start')
    end_date_str = request.GET.get('end')

    start_date = datetime.strptime(start_date_str, '%Y-%m-%d').date()
    end_date = datetime.strptime(end_date_str, '%Y-%m-%d').date()

    activities = TeleworkActivity.objects.filter(
        employee=employee,
        created_at__date__range=(start_date, end_date)
    ).order_by('created_at')

    activities_by_date = {}
    for activity in activities:
        d = activity.created_at.date()
        activities_by_date.setdefault(d, []).append(activity)

    letterhead_data = None
    config = SystemConfiguration.objects.filter(is_active=True).first()
    if config:
        image_field = (
                getattr(config, 'letterhead', None) or
                getattr(config, 'header_image', None) or
                getattr(config, 'logo', None) or
                getattr(config, 'institution_logo', None) or
                getattr(config, 'header_img', None)
        )
        if image_field and hasattr(image_field, 'path'):
            letterhead_data = image_field.path.replace('\\', '/')

    context = {
        'employee': employee,
        'start_date': start_date,
        'end_date': end_date,
        'activities_by_date': activities_by_date,
        'today': timezone.now().date(),
        'letterhead_data': letterhead_data,
    }

    template = get_template('biometric/reports/pdf_telework_report.html')
    html = template.render(context)
    result = BytesIO()
    pdf = pisa.pisaDocument(BytesIO(html.encode("UTF-8")), result)

    if not pdf.err:
        response = HttpResponse(result.getvalue(), content_type='application/pdf')
        filename = f"reporte_teletrabajo_{person.document_number}_{start_date_str}_a_{end_date_str}.pdf"
        response['Content-Disposition'] = f'inline; filename="{filename}"'
        return response

    return HttpResponse("Error al generar el PDF", status=500)


@login_required
def get_cv_stats_api(request, person_id):
    person = get_object_or_404(Person, pk=person_id)
    curriculum = getattr(person, 'curriculum', None)

    titles_count = curriculum.academic_titles.count() if curriculum else 0
    experiences_count = curriculum.work_experiences.count() if curriculum else 0
    courses_count = curriculum.trainings.count() if curriculum else 0
    experience_years, experience_months = 0, 0

    if curriculum:
        total_days = sum(
            ((date.today() if exp.is_current else (exp.end_date or date.today())) - exp.start_date).days
            for exp in curriculum.work_experiences.all() if exp.start_date
        )
        experience_years = total_days // 365
        experience_months = (total_days % 365) // 30

    return JsonResponse({
        'success': True,
        'titles_count': titles_count,
        'experiences_count': experiences_count,
        'courses_count': courses_count,
        'experience_text': f"{experience_years} años {experience_months} meses"
    })


@require_POST
def bulk_update_tab_visibility(request):
    tab_id = request.POST.get('tab_id')
    is_visible = request.POST.get('is_visible') == 'true'

    if not tab_id:
        return JsonResponse({'success': False, 'message': 'ID de pestaña no válido.'})

    try:
        with transaction.atomic():
            users = User.objects.filter(person__isnull=False)
            for u in users:
                EmployeeProfileVisibility.objects.update_or_create(user=u, tab_id=tab_id,
                                                                   defaults={'is_visible': is_visible})
        return JsonResponse({'success': True, 'message': f'Pestaña "{tab_id}" actualizada para todos los empleados.'})
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)}, status=500)


@method_decorator(require_POST, name='dispatch')
class UpdateProfileVisibilityView(LoginRequiredMixin, PermissionRequiredMixin, View):
    permission_required = 'person.change_person'

    def post(self, request, *args, **kwargs):
        try:
            data = json.loads(request.body)
            person_id = data.get('user_id')
            tab_id = data.get('tab_id')
            is_visible = data.get('is_visible')

            if not all([person_id, tab_id, isinstance(is_visible, bool)]):
                return JsonResponse({'success': False, 'message': 'Faltan datos o son incorrectos.'}, status=400)

            person = get_object_or_404(Person, pk=person_id)
            user = person.user
            if not user:
                return JsonResponse({'success': False, 'message': 'La persona no tiene un usuario asociado.'},
                                    status=400)

            visibility, created = EmployeeProfileVisibility.objects.update_or_create(
                user=user, tab_id=tab_id, defaults={'is_visible': is_visible}
            )
            return JsonResponse({'success': True, 'message': 'Visibilidad actualizada correctamente.'})
        except Exception as e:
            return JsonResponse({'success': False, 'message': str(e)}, status=500)


@login_required
def relocate_employee(request):
    try:
        person_id = request.POST.get('person_id')
        unit_id = request.POST.get('unit_id')
        if not person_id or not unit_id:
            return JsonResponse({'success': False, 'message': 'Faltan parámetros.'}, status=400)
        person = get_object_or_404(Person, pk=person_id)
        employee = getattr(person, 'employee_profile', None)
        if not employee:
            return JsonResponse({'success': False, 'message': 'No se encontró el perfil de empleado.'}, status=404)
        employee.area_id = unit_id
        employee.save()
        return JsonResponse({'success': True, 'message': 'Empleado reubicado correctamente.'})
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)}, status=500)


@login_required
def export_person_audit_excel(request, person_id):
    person = get_object_or_404(Person, pk=person_id)
    today = timezone.now().date()
    start_date_str = request.GET.get('start_date')
    end_date_str = request.GET.get('end_date')
    show_all = request.GET.get('show_all') == '1'

    logs_qs = PersonAuditLog.objects.filter(person=person).select_related('user').order_by('-created_at')

    if not show_all:
        if not start_date_str:
            start_date_str = today.isoformat()
        if not end_date_str:
            end_date_str = today.isoformat()
        try:
            start_d = datetime.strptime(start_date_str, '%Y-%m-%d').date()
            end_d = datetime.strptime(end_date_str, '%Y-%m-%d').date()
            logs_qs = logs_qs.filter(created_at__date__range=(start_d, end_d))
        except ValueError:
            logs_qs = logs_qs.filter(created_at__date=today)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Auditoría"

    headers = ["Fecha y Hora", "Usuario", "Movimiento / Detalle", "Dirección IP"]
    ws.append(headers)

    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for log in logs_qs:
        user_name = log.user.get_full_name() or log.user.username if log.user else "Sistema"
        action_text = log.action_detail or f"{log.get_action_display()} {log.get_section_display() or log.section or ''}".strip()
        created_str = log.created_at.strftime("%d/%m/%Y %H:%M:%S") if log.created_at else ""
        ws.append([created_str, user_name, action_text, log.ip_address or "—"])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    filename = f"auditoria_{person.document_number}_{timezone.now().strftime('%Y%m%d_%H%M')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response
